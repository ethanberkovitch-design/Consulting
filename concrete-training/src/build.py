"""Builds the "תורת הבטון" training deck on top of the corporate template.

Usage: python3 build.py <template.pptx> <assets_dir> <out.pptx>
The template's master (logo, line, slide number, title style) is kept as is; the
cover, section-divider and thank-you slides are cloned from the template's samples.
"""
import copy
import json
import re
import sys

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

TEMPLATE, ASSETS, OUT = sys.argv[1:4]

FONT = "Arial Unicode MS"  # mandated by the template
NAVY, BLUE, TITLE_BLUE, LIME, LIME_D = "003973", "0081C8", "0070C0", "C4D82E", "7E8F12"
GRAY, TEXT, MUTED, WHITE = "A3B1BA", "333F4A", "5F6B76", "FFFFFF"
TINT, TINT2, LIME_T = "EAF3FA", "F3F5F7", "F3F7D9"
RED, RED_T, GREEN, GREEN_T, AMBER, AMBER_T = "C62828", "FBE9E9", "2E7D32", "E8F3E9", "B26A00", "FFF2DC"
FOOTER = "סולל בונה תעשיות | אקדמיית טכנולוגיה ואיכות"

prs = Presentation(TEMPLATE)
TPL = list(prs.slides)  # 0 cover, 6 divider, 7 thanks
L_TITLE_ONLY = prs.slide_layouts[5]


def I(v):
    return Emu(int(v * 914400))


def rgb(h):
    return RGBColor.from_string(h)


RANGE = re.compile(r"(\d+(?:\.\d+)?)%-(\d+(?:\.\d+)?)%")


def bidi(t):
    """'70%-80%' flips to '80%-70%' inside RTL text; '70-80%' stays one LTR run."""
    return RANGE.sub(r"\1-\2%", t)


def rl(t):
    """Chart text is laid out LTR; a leading RLM keeps Hebrew labels in reading order."""
    return "\u200F" + t


# ---------------------------------------------------------------- primitives
def set_font(run, size=None, color=None, bold=None, italic=None):
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", FONT)
    rPr.set("lang", "he-IL")
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic


def _para_fmt(p, align, rtl=True, bullet=False, space_after=0, line=None, level_indent=0.25):
    pPr = p._p.get_or_add_pPr()
    pPr.set("algn", {"r": "r", "l": "l", "c": "ctr", "j": "just"}[align])
    pPr.set("rtl", "1" if rtl else "0")
    for tag in ("a:buNone", "a:buChar", "a:buFont"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    if bullet:
        pPr.set("marL", str(int(level_indent * 914400)))
        pPr.set("indent", str(int(-level_indent * 914400)))
        bf = etree.SubElement(pPr, qn("a:buFont"))
        bf.set("typeface", "Arial")
        bc = etree.SubElement(pPr, qn("a:buChar"))
        bc.set("char", "•")
    else:
        pPr.set("marL", "0")
        pPr.set("indent", "0")
        etree.SubElement(pPr, qn("a:buNone"))
    if space_after:
        p.space_after = Pt(space_after)
    if line:
        p.line_spacing = line


def text(s, x, y, w, h, content, size=16, color=TEXT, bold=False, align="r", anchor="t",
         bullet=False, space_after=4, line=None, rtl=True, name=None):
    """content: str | list of paragraphs; paragraph: str | list of (text, opts) runs | dict."""
    tb = s.shapes.add_textbox(I(x), I(y), I(w), I(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    tb.text_frame._txBody.bodyPr.set("rtlCol", "1" if rtl else "0")
    paras = content if isinstance(content, list) else [content]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        opts = {}
        if isinstance(para, dict):
            opts = para
            para = para["t"]
        runs = para if isinstance(para, list) else [(para, {})]
        _para_fmt(p, opts.get("align", align), opts.get("rtl", rtl), opts.get("bullet", bullet),
                  opts.get("space_after", space_after), opts.get("line", line))
        for rt, ro in runs:
            r = p.add_run()
            r.text = bidi(rt)
            set_font(r, ro.get("size", opts.get("size", size)), ro.get("color", opts.get("color", color)),
                     ro.get("bold", opts.get("bold", bold)), ro.get("italic"))
    return tb


def _strip_style(shape):
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def box(s, x, y, w, h, fill=TINT, line=None, lw=1.0, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08,
        name=None, rotation=0):
    sh = s.shapes.add_shape(shape, I(x), I(y), I(w), I(h))
    _strip_style(sh)
    if fill:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    else:
        sh.fill.background()
    if line:
        sh.line.color.rgb = rgb(line)
        sh.line.width = Pt(lw)
    else:
        sh.line.fill.background()
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    if rotation:
        sh.rotation = rotation
    if name:
        sh.name = name
    return sh


def shape_text(sh, content, size=16, color=WHITE, bold=True, align="c", anchor="m", margin=0.05):
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = I(margin)
    tf.margin_top = tf.margin_bottom = I(0.02)
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    paras = content if isinstance(content, list) else [content]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _para_fmt(p, align)
        runs = para if isinstance(para, list) else [(para, {})]
        for rt, ro in runs:
            r = p.add_run()
            r.text = bidi(rt)
            set_font(r, ro.get("size", size), ro.get("color", color), ro.get("bold", bold))
    return sh


def line(s, x1, y1, x2, y2, color=BLUE, w=2.0, head=True, dash=False, tail=False):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, I(x1), I(y1), I(x2), I(y2))
    _strip_style(c)
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(w)
    ln = c.line._get_or_add_ln()
    if dash:
        pd = etree.SubElement(ln, qn("a:prstDash"))
        pd.set("val", "dash")
    if tail:
        e = etree.SubElement(ln, qn("a:headEnd"))
        e.set("type", "triangle")
    if head:
        e = etree.SubElement(ln, qn("a:tailEnd"))
        e.set("type", "triangle")
    return c


def pic(s, name, x, y, w=None, h=None):
    path = f"{ASSETS}/{name}.png"
    iw, ih = Image.open(path).size
    if w and not h:
        h = w * ih / iw
    elif h and not w:
        w = h * iw / ih
    return s.shapes.add_picture(path, I(x), I(y), I(w), I(h))


def fit(name, x, y, w, h):
    """Largest size of an asset that fits the box, centred in it."""
    iw, ih = Image.open(f"{ASSETS}/{name}.png").size
    sc = min(w / iw, h / ih)
    ww, hh = iw * sc, ih * sc
    return x + (w - ww) / 2, y + (h - hh) / 2, ww, hh


def pic_fit(s, name, x, y, w, h):
    fx, fy, fw, fh = fit(name, x, y, w, h)
    return pic(s, name, fx, fy, fw, fh)


def icon(s, name, cx, cy, d, fill=BLUE, color="w", ring=None):
    c = box(s, cx - d / 2, cy - d / 2, d, d, fill=fill, shape=MSO_SHAPE.OVAL, line=ring, lw=2.5)
    g = d * 0.52
    pic(s, f"{name}_{color}", cx - g / 2, cy - g / 2, g, g)
    return c


def glyph(s, name, x, y, d, color="b"):
    return pic(s, f"{name}_{color}", x, y, d, d)


def badge(s, n, cx, cy, d=0.5, fill=NAVY, color=WHITE, size=16):
    b = box(s, cx - d / 2, cy - d / 2, d, d, fill=fill, shape=MSO_SHAPE.OVAL)
    shape_text(b, str(n), size=size, color=color, margin=0)
    return b


def card(s, x, y, w, h, title, body, ic=None, fill=TINT, tsize=18, bsize=15, title_color=NAVY,
         icon_fill=BLUE, bullet=False):
    box(s, x, y, w, h, fill=fill)
    pad = 0.2
    tx_w = w - 2 * pad
    ty = y + pad
    if ic:
        icon(s, ic, x + w - pad - 0.3, y + pad + 0.3, 0.6, fill=icon_fill)
        text(s, x + pad, ty + 0.08, tx_w - 0.75, 0.5, title, size=tsize, color=title_color, bold=True, anchor="m")
        ty += 0.75
    else:
        text(s, x + pad, ty, tx_w, 0.45, title, size=tsize, color=title_color, bold=True)
        ty += 0.5
    if body:
        text(s, x + pad, ty, tx_w, y + h - ty - pad, body, size=bsize, color=TEXT, bullet=bullet, space_after=3)


def message(s, content, y=5.78, x=0.9, w=8.95, h=0.7, ic="FaLightbulb"):
    """Key-message strip; kept clear of the template's logo at the bottom right."""
    box(s, x, y, w, h, fill=NAVY)
    icon(s, ic, x + w - 0.42, y + h / 2, 0.5, fill=LIME, color="n")
    text(s, x + 0.25, y, w - 1.0, h, content, size=17, color=WHITE, bold=True, anchor="m")


def table(s, x, y, w, rows, colw, size=13, header_fill=BLUE, row_h=0.42, header_h=0.5,
          zebra=(WHITE, TINT2), align=None, bold_first=True, cell_fill=None, header_size=None):
    """rows[0] is the header; columns are given right-to-left (logical RTL order)."""
    n, m = len(rows), len(rows[0])
    rows_ltr = [list(reversed(r)) for r in rows]
    colw_ltr = list(reversed(colw))
    align_ltr = list(reversed(align)) if align else ["r"] * m
    total_h = header_h + row_h * (n - 1)
    gf = s.shapes.add_table(n, m, I(x), I(y), I(w), I(total_h))
    tbl = gf.table
    tbl.first_row = True
    tbl.horz_banding = False
    scale = w / sum(colw_ltr)
    for j, cw in enumerate(colw_ltr):
        tbl.columns[j].width = I(cw * scale)
    for i in range(n):
        tbl.rows[i].height = I(header_h if i == 0 else row_h)
        for j in range(m):
            c = tbl.cell(i, j)
            c.margin_left = c.margin_right = I(0.08)
            c.margin_top = c.margin_bottom = I(0.04)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            fill = header_fill if i == 0 else zebra[(i - 1) % 2]
            if cell_fill and i > 0:
                fill = cell_fill(i, m - 1 - j) or fill
            c.fill.solid()
            c.fill.fore_color.rgb = rgb(fill)
            val = rows_ltr[i][j]
            tf = c.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            _para_fmt(p, "c" if i == 0 else align_ltr[j])
            r = p.add_run()
            r.text = bidi(str(val))
            is_first = (m - 1 - j) == 0
            set_font(r, (header_size or size) if i == 0 else size, WHITE if i == 0 else TEXT,
                     i == 0 or (bold_first and is_first))
    return gf


def style_chart(ch, size=12):
    ch.font.size = Pt(size)
    ch.font.name = FONT
    ch.font.color.rgb = rgb(TEXT)


def axis_fmt(ax, title=None, size=12, gridlines=True, fmt=None):
    ax.has_major_gridlines = gridlines
    if gridlines:
        ax.major_gridlines.format.line.color.rgb = rgb("DDE3E8")
        ax.major_gridlines.format.line.width = Pt(0.75)
    ax.format.line.color.rgb = rgb("B8C2CA")
    ax.tick_labels.font.size = Pt(size)
    ax.tick_labels.font.name = FONT
    ax.tick_labels.font.color.rgb = rgb(MUTED)
    if fmt:
        ax.tick_labels.number_format = fmt
        ax.tick_labels.number_format_is_linked = False
    if title:
        ax.has_title = True
        ax.axis_title.text_frame.text = title
        for p in ax.axis_title.text_frame.paragraphs:
            for r in p.runs:
                set_font(r, size, MUTED, False)


def series_line(ser, color, w=2.75, marker=True, dash=False, smooth=False):
    ser.format.line.color.rgb = rgb(color)
    ser.format.line.width = Pt(w)
    ser.smooth = smooth
    if dash:
        ln = ser.format.line._get_or_add_ln()
        pd = etree.SubElement(ln, qn("a:prstDash"))
        pd.set("val", "dash")
    if marker:
        ser.marker.style = XL_MARKER_STYLE.CIRCLE
        ser.marker.size = 7
        ser.marker.format.fill.solid()
        ser.marker.format.fill.fore_color.rgb = rgb(color)
        ser.marker.format.line.color.rgb = rgb(WHITE)
    else:
        ser.marker.style = XL_MARKER_STYLE.NONE


def verify(s, x, y, w=1.75, t="לאימות מול התקן"):
    """Marks values taken from common practice rather than the standard's text."""
    r = box(s, x, y, w, 0.32, fill=AMBER_T, line=AMBER, lw=1, radius=0.5)
    shape_text(r, t, size=11, color=AMBER, margin=0.02)


def caption(s, x, y, w, t, align="r"):
    text(s, x, y, w, 0.3, t, size=11, color=MUTED, align=align)


# ---------------------------------------------------------------- slide factories
SECTION_LOG = []


def content_slide(title):
    s = prs.slides.add_slide(L_TITLE_ONLY)
    s.shapes.title.text = title
    for p in s.shapes.title.text_frame.paragraphs:
        for r in p.runs:
            set_font(r)
    return s


def clone(src):
    s = prs.slides.add_slide(src.slide_layout)
    for shp in list(s.shapes):
        shp._element.getparent().remove(shp._element)
    rid_map = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype.endswith("/image"):
            rid_map[rid] = s.part.relate_to(rel._target, rel.reltype)
    tree = s.shapes._spTree
    for el in src.shapes._spTree.iterchildren():
        if el.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            continue
        new = copy.deepcopy(el)
        for node in new.iter():
            for attr in (qn("r:embed"), qn("r:id"), qn("r:link")):
                if node.get(attr) in rid_map:
                    node.set(attr, rid_map[node.get(attr)])
        tree.append(new)
    return s


def find(s, name):
    return next(sh for sh in s.shapes if sh.name == name)


def retext(sh, content, size=None, color=None, align=None, wrap=True):
    """Replace a cloned text box's text, keeping its first run's formatting."""
    tf = sh.text_frame
    first_p = tf.paragraphs[0]._p
    rpr = copy.deepcopy(first_p.find(qn("a:r")).find(qn("a:rPr")))
    ppr = first_p.find(qn("a:pPr"))
    ppr = copy.deepcopy(ppr) if ppr is not None else None
    body = tf._txBody
    for p in body.findall(qn("a:p")):
        body.remove(p)
    for line_ in content if isinstance(content, list) else [content]:
        p = etree.SubElement(body, qn("a:p"))
        if ppr is not None:
            p.append(copy.deepcopy(ppr))
        r = etree.SubElement(p, qn("a:r"))
        rp = copy.deepcopy(rpr)
        if size:
            rp.set("sz", str(int(size * 100)))
        r.append(rp)
        t = etree.SubElement(r, qn("a:t"))
        t.text = line_
    bp = body.find(qn("a:bodyPr"))
    if wrap:
        bp.set("wrap", "square")
        for el in list(bp):
            if el.tag in (qn("a:spAutoFit"), qn("a:normAutofit")):
                bp.remove(el)
    for p in tf.paragraphs:
        pPr = p._p.get_or_add_pPr()
        if align:
            pPr.set("algn", align)
        pPr.set("rtl", "1")
        if color:
            for r in p.runs:
                r.font.color.rgb = rgb(color)


def divider(num, title, topics):
    s = clone(TPL[6])
    tb = find(s, "TextBox 7")
    tb.left, tb.top, tb.width, tb.height = I(7.35), I(2.75), I(5.4), I(1.0)
    retext(tb, title, size=36, align="ctr")
    tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    text(s, 7.35, 2.05, 5.4, 0.5, f"פרק {num}" if num else "", size=20, color=WHITE, align="c")
    text(s, 7.75, 3.95, 4.6, 2.2, " · ".join(topics), size=16, color=WHITE, align="c", line=1.15)
    return s


# ---------------------------------------------------------------- footer on the content layout
def add_layout_footer():
    tmp = prs.slides.add_slide(L_TITLE_ONLY)
    tb = text(tmp, 0.95, 6.82, 4.4, 0.3, FOOTER, size=10, color=MUTED, align="l", name="Footer academy")
    el = tb._element
    L_TITLE_ONLY.shapes._spTree.append(el)
    # drop the temporary slide
    sldIdLst = prs.slides._sldIdLst
    rid = sldIdLst[-1].get(qn("r:id"))
    prs.part.drop_rel(rid)
    sldIdLst.remove(sldIdLst[-1])


add_layout_footer()

# ================================================================ 1. Cover
s = clone(TPL[0])
t = find(s, "TextBox 4")
t.left, t.top, t.width, t.height = I(1.67), I(3.78), I(10.0), I(0.8)
retext(t, "תורת הבטון", size=40, align="ctr")
d = find(s, "מציין מיקום תוכן 2")
d.left, d.width = I(0.32), I(2.6)
retext(d, "אוקטובר 2026", size=12)
for p in d.text_frame.paragraphs:
    _para_fmt(p, "l")
    for r in p.runs:
        set_font(r, 12, MUTED)
text(s, 1.67, 4.6, 10.0, 0.45, "יסודות, איכות ומצוינות מקצועית", size=22, color=NAVY, align="c")
text(s, 1.67, 5.06, 10.0, 0.35, "Concrete Technology Fundamentals  |  אקדמיית טכנולוגיה ואיכות", size=13,
     color=MUTED, align="c")
pic(s, "skyline", 2.05, 5.55, w=9.2)

# ================================================================ 2. Objectives
s = content_slide("מטרות ההדרכה")
cx, cy = 3.1, 3.55
for d_, col in ((4.3, "DCE6F1"), (3.3, "A9C6E3"), (2.3, BLUE), (1.2, NAVY)):
    box(s, cx - d_ / 2, cy - d_ / 2, d_, d_, fill=col, shape=MSO_SHAPE.OVAL)
pic(s, "FaBullseye_w", cx - 0.32, cy - 0.32, 0.64, 0.64)
line(s, 5.6, 1.65, cx + 0.25, cy - 0.2, color=LIME_D, w=4)
text(s, 0.9, 5.95, 4.4, 0.6, "מטרה: משתתף שמבין את החומר — ומונע כשלים לפני שהם קורים", size=14,
     color=MUTED, align="c")
objs = [
    ("FaFlask", "להבין כיצד נוצר בטון", "מרכיבים, הידרציה ומבנה"),
    ("FaCogs", "להבין את גורמי האיכות", "W/C, חומרי גלם, אשפרה"),
    ("FaExclamationTriangle", "לזהות כשלים", "סגרגציה, דימום, קיני חצץ, סדיקה"),
    ("FaClipboardCheck", "לבצע בקרת איכות", "בדיקות טרי, קשוח וחומרי גלם"),
    ("FaBook", "להכיר תקנים ודרישות", "ת\"י 118, ת\"י 26, ת\"י 3, ת\"י 1"),
]
for i, (ic, t1, t2) in enumerate(objs):
    y = 1.2 + i * 0.98
    icon(s, ic, 12.2, y + 0.4, 0.7)
    text(s, 6.2, y + 0.02, 5.5, 0.42, t1, size=19, color=NAVY, bold=True)
    text(s, 6.2, y + 0.45, 5.5, 0.35, t2, size=15, color=MUTED)

# ================================================================ 3. Value chain
s = content_slide("שרשרת הערך של הבטון")
chain = [
    ("FaMountain", "מחצבה", "סוג הסלע וניקיונו"),
    ("FaFilter", "מפעל אגרגטים", "גריסה, ניפוי, שטיפה"),
    ("FaWarehouse", "צמנט", "סוג, חוזק, אחידות"),
    ("FaIndustry", "מפעל בטון", "תכן תערובת, שקילה, ערבוב"),
    ("FaTruck", "הובלה", "זמן, טמפרטורה, ללא תוספת מים"),
    ("FaHardHat", "אתר", "יציקה, ריטוט, אשפרה"),
    ("FaBuilding", "לקוח", "מבנה עמיד לעשרות שנים"),
]
step_w, gap = 1.45, 0.25
x0 = 12.55 - step_w
for i, (ic, t1, t2) in enumerate(chain):
    x = x0 - i * (step_w + gap)
    fill = NAVY if i in (3,) else BLUE
    icon(s, ic, x + step_w / 2, 1.95, 1.0, fill=fill)
    text(s, x - 0.05, 2.6, step_w + 0.1, 0.42, t1, size=17, color=NAVY, bold=True, align="c")
    text(s, x - 0.05, 3.02, step_w + 0.1, 0.8, t2, size=13, color=MUTED, align="c")
    if i < len(chain) - 1:
        line(s, x - 0.02, 1.95, x - gap + 0.02, 1.95, color=GRAY, w=2.5)
box(s, 0.9, 4.05, 11.65, 1.45, fill=TINT2)
pts = [("FaSearch", "כל חוליה משפיעה", "אגרגט מלוכלך במחצבה = בעיית חוזק ועבידות במפעל"),
       ("FaLink", "האחריות משותפת", "מספק חומרי הגלם ועד הצוות שמאשפר באתר"),
       ("FaChartLine", "בקרה לאורך השרשרת", "בדיקה בכל מעבר — לא רק בקוביות של 28 יום")]
for i, (ic, t1, t2) in enumerate(pts):
    x = 8.75 - i * 3.9
    icon(s, ic, x + 3.45, 4.5, 0.5, fill=NAVY)
    text(s, x, 4.3, 3.1, 0.4, t1, size=16, color=NAVY, bold=True)
    text(s, x, 4.7, 3.1, 0.75, t2, size=13, color=TEXT)
message(s, "איכות הבטון נקבעת לאורך כל שרשרת הערך — חוליה חלשה אחת מספיקה לכשל")

# ================================================================ Section 1
divider(1, "מרכיבי הבטון", ["אגרגטים", "צמנט", "מים", "מוספים ותוספים"])

# ---------------------------------------------------------------- 4. What is concrete
s = content_slide("מהו בטון?")
text(s, 5.6, 1.15, 6.95, 0.75, "חומר מרוכב: שלד של אגרגטים שמשחת צמנט ומים מלכדת לגוף אחד מוצק",
     size=17, color=MUTED)
cd = CategoryChartData()
parts = [("אגרגט גס", 42, NAVY), ("אגרגט דק", 28, BLUE), ("מים", 17, "7FC3EA"), ("צמנט", 11, GRAY),
         ("אוויר", 2, LIME)]
cd.categories = [p[0] for p in parts]
cd.add_series("נפח", [p[1] for p in parts])
gf = s.shapes.add_chart(XL_CHART_TYPE.DOUGHNUT, I(0.9), I(1.15), I(4.6), I(4.2), cd)
ch = gf.chart
style_chart(ch, 13)
ch.has_legend = False
ch.has_title = False
plot = ch.plots[0]
for i, p in enumerate(parts):
    pt = plot.series[0].points[i]
    pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = rgb(p[2])
    pt.format.line.color.rgb = rgb(WHITE)
plot.has_data_labels = True
dl = plot.data_labels
dl.number_format = '0"%"'
dl.number_format_is_linked = False
dl.font.size = Pt(13)
dl.font.bold = True
dl.font.color.rgb = rgb(WHITE)
hole = ch._chartSpace.find(".//" + qn("c:holeSize"))
if hole is not None:
    hole.set("val", "55")
text(s, 2.15, 2.85, 2.1, 0.7, [{"t": "כ-70%", "size": 26, "bold": True, "color": NAVY},
                              {"t": "אגרגטים", "size": 14, "color": MUTED}], align="c")
caption(s, 0.9, 5.35, 4.6, "חלוקת נפח טיפוסית בבטון רגיל (להמחשה)", align="c")
comps = [
    ("FaCubes", NAVY, "אגרגטים", "חצץ וחול — השלד הנושא של הבטון"),
    ("FaWarehouse", GRAY, "צמנט", "החומר הקושר: מגיב עם מים ויוצר את הדבק"),
    ("FaTint", "7FC3EA", "מים", "מניעים את ההידרציה ומעניקים עבידות"),
    ("FaVial", BLUE, "מוספים", "כימיקלים במינון נמוך שמשנים התנהגות"),
    ("FaRecycle", LIME, "תוספים", "חומרים מינרליים פעילים במקום חלק מהצמנט"),
]
for i, (ic, col, t1, t2) in enumerate(comps):
    y = 1.95 + i * 0.72
    icon(s, ic, 12.2, y + 0.3, 0.58, fill=col, color="n" if col in (LIME, "7FC3EA", GRAY) else "w")
    text(s, 5.9, y + 0.02, 5.85, 0.6, [[(t1 + "  ", {"bold": True, "color": NAVY, "size": 17}),
                                        (t2, {"size": 15})]], anchor="m")
message(s, "בטון טוב = חומרי גלם טובים × תכן נכון × ביצוע נכון", y=5.8, h=0.62)

# ---------------------------------------------------------------- 5. Aggregates
s = content_slide("תפקיד האגרגטים")
pic_fit(s, "aggregates", 0.9, 1.2, 6.6, 4.55)
caption(s, 0.9, 5.8, 6.6, "חתך בטון: האגרגטים (מסומנים) הם רוב הנפח, ומשחת הצמנט ממלאת את החללים", align="c")
text(s, 7.9, 1.15, 4.65, 1.2, [{"t": "70%-80%", "size": 54, "bold": True, "color": BLUE},
                               {"t": "מנפח הבטון", "size": 18, "color": MUTED}], align="r")
pts = [("FaWeightHanging", "שלד נושא", "תורמים לחוזק ולמודול האלסטיות; אגרגט חזק ונקי = בטון חזק"),
       ("FaShieldAlt", "עמידות ויציבות נפחית", "מקטינים התכווצות, זחילה וחום הידרציה"),
       ("FaCoins", "כלכליות", "זולים מהצמנט — גרדציה טובה חוסכת משחה")]
for i, (ic, t1, t2) in enumerate(pts):
    y = 2.6 + i * 1.08
    icon(s, ic, 12.25, y + 0.33, 0.6)
    text(s, 7.9, y, 3.95, 0.38, t1, size=18, color=NAVY, bold=True)
    text(s, 7.9, y + 0.4, 3.95, 0.65, t2, size=14)

# ---------------------------------------------------------------- 6. Cement
s = content_slide("תפקיד הצמנט")
pic_fit(s, "hydration", 0.9, 1.35, 7.4, 2.4)
for i, (t1, t2) in enumerate([("גרגרי צמנט במים", "דקות 0"), ("תחילת הידרציה", "שעות"), ("רשת C-S-H מלכדת", "ימים")]):
    x = 0.9 + i * 2.53
    text(s, x, 3.85, 2.3, 0.35, t1, size=15, color=NAVY, bold=True, align="c")
    text(s, x, 4.2, 2.3, 0.3, t2, size=13, color=MUTED, align="c")
box(s, 0.9, 4.75, 7.4, 0.85, fill=TINT2)
text(s, 1.1, 4.8, 7.0, 0.75, [[("מרכיבי הקלינקר:  ", {"bold": True, "color": NAVY}), ("C₃S, C₂S, C₃A, C₄AF", {})],
                              [("סוגי צמנט:  ", {"bold": True, "color": NAVY}), ("CEM I עד CEM V, לפי ת\"י 1", {})]],
     size=14, anchor="m", space_after=2)
pts = [("FaLink", "החומר הקושר", "המרכיב היחיד שהופך את התערובת לגוף מוצק"),
       ("FaLayerGroup", "יוצר את משחת הצמנט", "צמנט + מים עוטפים את האגרגטים וממלאים חללים"),
       ("FaChartLine", "מפתח חוזק", "התגובה נמשכת שבועות וחודשים כל עוד יש לחות")]
for i, (ic, t1, t2) in enumerate(pts):
    y = 1.3 + i * 1.2
    icon(s, ic, 12.25, y + 0.35, 0.62)
    text(s, 8.7, y, 3.15, 0.4, t1, size=18, color=NAVY, bold=True)
    text(s, 8.7, y + 0.42, 3.15, 0.75, t2, size=14)
message(s, "צמנט אינו \"דבק\" שמתייבש — הוא מגיב כימית עם מים ומתקשה", y=5.85, h=0.62)

# ---------------------------------------------------------------- 7. Water
s = content_slide("תפקיד המים")
text(s, 0.9, 1.2, 7.5, 0.4, "יחס מים-צמנט (W/C) — סקאלת השפעה", size=17, color=NAVY, bold=True)
zones = [(0.30, 0.40, NAVY, "חוזק גבוה מאוד", "דורש סופר-פלסטיצייזר"),
         (0.40, 0.50, BLUE, "בטון מבני איכותי", "עמידות טובה"),
         (0.50, 0.60, "7FC3EA", "בטון רגיל", "רגיש לאשפרה"),
         (0.60, 0.75, GRAY, "בטון חלש וחדיר", "סיכון לעמידות")]
sx, sw = 0.9, 7.5
lo, hi = 0.30, 0.75
for a, b_, col, t1, t2 in zones:
    # higher W/C to the left so the scale reads right-to-left like the slide
    x = sx + (hi - b_) / (hi - lo) * sw
    w = (b_ - a) / (hi - lo) * sw
    r = box(s, x, 1.8, w - 0.04, 0.65, fill=col, shape=MSO_SHAPE.RECTANGLE)
    shape_text(r, t1, size=13, color=WHITE if col in (NAVY, BLUE) else NAVY)
    text(s, x, 2.5, w - 0.04, 0.3, t2, size=12, color=MUTED, align="c")
for v in (0.30, 0.40, 0.50, 0.60, 0.75):
    x = sx + (hi - v) / (hi - lo) * sw
    text(s, x - 0.3, 2.85, 0.6, 0.3, f"{v:.2f}", size=13, color=TEXT, align="c", rtl=False)
line(s, sx + 7.3, 3.35, sx + 0.2, 3.35, color=MUTED, w=1.5)
text(s, sx, 3.42, sw, 0.3, "חוזק ועמידות יורדים  ←  יותר מים", size=13, color=MUTED, align="c")
box(s, 0.9, 4.0, 3.6, 1.55, fill=TINT)
text(s, 1.05, 4.08, 3.3, 1.45, [{"t": "0.25", "size": 32, "bold": True, "color": BLUE},
                                {"t": "W/C שנדרש כימית להידרציה מלאה (בקירוב)", "size": 14}], align="c")
box(s, 4.8, 4.0, 3.6, 1.55, fill=TINT)
text(s, 4.95, 4.08, 3.3, 1.45, [{"t": "0.45-0.55", "size": 32, "bold": True, "color": NAVY},
                                {"t": "W/C טיפוסי בבטון מבני — העודף הוא \"מי עבידות\"", "size": 14}], align="c")
pts = [("FaFlask", "יצירת הידרציה", "המים מגיבים עם הצמנט ויוצרים את מוצרי ההידרציה"),
       ("FaHandPaper", "השפעה על עבידות", "מאפשרים שאיבה, יציקה וריטוט"),
       ("FaArrowDown", "השפעה על חוזק", "כל ליטר עודף משאיר נקבוביות ומחליש")]
for i, (ic, t1, t2) in enumerate(pts):
    y = 1.3 + i * 1.3
    icon(s, ic, 12.25, y + 0.35, 0.62)
    text(s, 8.8, y, 3.05, 0.4, t1, size=18, color=NAVY, bold=True)
    text(s, 8.8, y + 0.42, 3.05, 0.85, t2, size=14)
message(s, "המים הנוספים מעבר לנדרש כימית — הם שמשאירים את הנקבוביות", y=5.85, h=0.62)

# ---------------------------------------------------------------- 8. Admixtures vs additions
s = content_slide("מוספים ותוספים")
box(s, 6.85, 1.15, 5.7, 0.8, fill=TINT)
icon(s, "FaVial", 12.15, 1.55, 0.55)
text(s, 7.0, 1.2, 4.75, 0.7, [[("מוספים (Admixtures): ", {"bold": True, "color": NAVY}),
                               ("כימיקלים, מינון קטן מ-5% ממשקל הצמנט", {})]], size=14, anchor="m")
box(s, 0.9, 1.15, 5.7, 0.8, fill=LIME_T)
icon(s, "FaRecycle", 6.2, 1.55, 0.55, fill=LIME, color="n")
text(s, 1.05, 1.2, 4.75, 0.7, [[("תוספים (Additions): ", {"bold": True, "color": NAVY}),
                               ("חומרים מינרליים פעילים, מחליפים חלק מהצמנט", {})]], size=14, anchor="m")
rows = [["סוג", "חומר", "פעולה עיקרית", "מתי משתמשים"],
        ["מוסף", "סופר-פלסטיצייזר", "מפזר את גרגרי הצמנט; מפחית מים ב-12%-30%", "חוזק גבוה, בטון זורם, W/C נמוך"],
        ["מוסף", "מעכב", "דוחה את תחילת ההתקשרות", "קיץ, הובלה ארוכה, יציקות רציפות"],
        ["מוסף", "מאיץ", "מזרז התקשרות והתפתחות חוזק מוקדם", "חורף, טרום, פירוק תבניות מוקדם"],
        ["תוסף", "אפר פחם (Fly Ash)", "תגובה פוצולנית; עבידות ועמידות", "יציקות מסיביות, עמידות לכלורידים"],
        ["תוסף", "מיקרוסיליקה", "חלקיקים זעירים שממלאים חללים", "בטון בחוזק גבוה, סביבה ימית"],
        ["תוסף", "סיגים (GGBS)", "חום הידרציה נמוך, עמידות כימית", "מבנים ימיים, סולפטים, מסיבי"]]
table(s, 0.9, 2.15, 11.65, rows, [1.0, 2.2, 4.4, 4.05], size=13, row_h=0.52,
      cell_fill=lambda i, j: (TINT if i <= 3 else LIME_T) if j == 0 else None)

# ================================================================ Section 2
divider(2, "הידרציה ומבנה", ["תהליך ההידרציה", "מיקרו-מבנה", "גורמים משפיעים", "התפתחות המבנה"])

# ---------------------------------------------------------------- 9. Hydration
s = content_slide("מהי הידרציה?")
steps = [("FaTint", "מים + צמנט", "מגע ראשוני, תחילת פירוק המינרלים"),
         ("FaBolt", "תגובה כימית", "אקסותרמית — משחררת חום"),
         ("FaLayerGroup", "C-S-H", "סיליקט סידן מוהדר: \"הדבק\" של הבטון"),
         ("FaChartLine", "התפתחות חוזק", "הרשת מתעבה והחללים מתמלאים")]
for i, (ic, t1, t2) in enumerate(steps):
    x = 9.75 - i * 2.95
    box(s, x, 1.3, 2.75, 2.7, fill=NAVY if i == 2 else TINT)
    icon(s, ic, x + 1.375, 1.95, 0.85, fill=LIME if i == 2 else BLUE, color="n" if i == 2 else "w")
    text(s, x + 0.15, 2.55, 2.45, 0.45, t1, size=20, bold=True, color=WHITE if i == 2 else NAVY, align="c")
    text(s, x + 0.15, 3.05, 2.45, 0.85, t2, size=14, color=WHITE if i == 2 else TEXT, align="c")
    if i < 3:
        line(s, x - 0.02, 2.65, x - 0.18, 2.65, color=BLUE, w=3)
box(s, 0.9, 4.3, 11.65, 1.2, fill=TINT2)
text(s, 1.1, 4.35, 5.5, 1.1, [{"t": "C₃S / C₂S + H₂O  →  C-S-H + CH", "size": 22, "bold": True, "color": NAVY,
                               "align": "c"},
                              {"t": "תוצר משני: סידן הידרוקסידי — שומר על pH גבוה ומגן על הזיון", "size": 13,
                               "align": "c"}], anchor="m", rtl=True)
text(s, 6.9, 4.4, 5.45, 1.05, [{"t": "50%-60% מנפח המשחה המוקשה הוא C-S-H", "bullet": True},
                               {"t": "ההידרציה צורכת מים — בלי לחות היא נעצרת", "bullet": True}], size=15)
message(s, "אין הידרציה — אין חוזק. לכן האשפרה היא חלק מהתהליך, לא תוספת", y=5.8, h=0.65)

# ---------------------------------------------------------------- 10. Microstructure
s = content_slide("מיקרו-מבנה של בטון")
anch = json.load(open(f"{ASSETS}/micro_anchors.json"))
mx, my, mw = 0.9, 1.15, 7.4
scale = mw / 1000
pic(s, "micro", mx, my, w=mw)
labels = [("אגרגט", "Aggregate", "השלד: חזק וצפוף"),
          ("ITZ — אזור המגע", "Interfacial Transition Zone", "10-50 מיקרון, נקבובי יותר — החוליה החלשה"),
          ("משחת צמנט", "Cement Paste", "C-S-H, CH ונקבוביות קפילריות")]
cols = [NAVY, LIME_D, TEXT]
for (t1, en, t2), ay, col in zip(labels, anch, cols):
    y = my + ay * scale
    text(s, 8.45, y - 0.5, 4.1, 0.4, [[(t1, {"bold": True, "color": col, "size": 18}),
                                       ("   " + en, {"color": MUTED, "size": 12})]])
    text(s, 8.45, y - 0.1, 4.1, 0.45, t2, size=14)
message(s, "שיפור ה-ITZ (W/C נמוך, מיקרוסיליקה, אגרגט נקי) = בטון חזק ועמיד יותר", y=5.85, h=0.62)

# ---------------------------------------------------------------- 11. Factors on hydration
s = content_slide("גורמים המשפיעים על הידרציה")
cx, cy = 6.72, 3.45
box(s, cx - 1.15, cy - 1.15, 2.3, 2.3, fill=NAVY, shape=MSO_SHAPE.OVAL)
text(s, cx - 1.1, cy - 0.4, 2.2, 0.8, "הידרציה", size=24, bold=True, color=WHITE, align="c", anchor="m")
facs = [("FaThermometerHalf", "טמפרטורה", "חום מאיץ את התגובה המוקדמת אך פוגע בחוזק הסופי; מתחת ל-5°C היא כמעט נעצרת", 9.0, 1.2),
        ("FaTint", "יחס W/C", "קובע כמה מים זמינים וכמה חללים יישארו", 0.9, 1.2),
        ("FaWarehouse", "סוג הצמנט", "דקות טחינה והרכב: CEM I מהיר, CEM III עם סיגים איטי", 9.0, 3.85),
        ("FaClock", "זמן", "ההידרציה נמשכת חודשים — כל עוד יש לחות", 0.9, 3.85)]
for ic, t1, t2, x, y in facs:
    box(s, x, y, 3.55, 1.9, fill=TINT)
    icon(s, ic, x + 3.55 - 0.5, y + 0.5, 0.65)
    text(s, x + 0.2, y + 0.25, 2.4, 0.45, t1, size=19, bold=True, color=NAVY)
    text(s, x + 0.2, y + 0.85, 3.15, 1.0, t2, size=14)
    ex = x if x > cx else x + 3.55
    line(s, cx + (1.1 if x > cx else -1.1), cy + (-0.5 if y < 3 else 0.5), ex, y + 0.95, color=GRAY, w=2,
         head=False)

# ---------------------------------------------------------------- 12. Structure development
s = content_slide("התפתחות המבנה: מבטון טרי לבטון קשוח")
stages = [("1", "פלסטי", "0-2 שעות", "ניתן להובלה, שאיבה, יציקה וריטוט", "7FC3EA"),
          ("2", "תחילת התקשרות", "2-4 שעות", "סוף חלון הריטוט; אסור להוסיף מים או לערבב מחדש", BLUE),
          ("3", "התקשות", "4-10 שעות", "סוף התקשרות; מתחילה עבודת החלקה והגנה מהתאדות", "1F5A94"),
          ("4", "קשיח", "מ-24 שעות", "צובר חוזק; פירוק תבניות לפי חוזק בפועל", NAVY)]
tw = 11.65 / 4
for i, (n, t1, tm, t2, col) in enumerate(stages):
    x = 12.55 - (i + 1) * tw
    r = box(s, x + 0.02, 2.0, tw - 0.04, 0.75, fill=col, shape=MSO_SHAPE.RECTANGLE)
    shape_text(r, t1, size=19)
    badge(s, n, x + tw - 0.35, 1.55, 0.55, fill=LIME, color=NAVY)
    text(s, x + 0.1, 2.85, tw - 0.2, 0.35, tm, size=15, color=BLUE, bold=True, align="c")
    text(s, x + 0.15, 3.25, tw - 0.3, 1.1, t2, size=14, align="c")
line(s, 12.55, 4.45, 0.95, 4.45, color=MUTED, w=2)
text(s, 0.9, 4.52, 3.5, 0.3, "זמן", size=13, color=MUTED, align="l")
caption(s, 0.9, 4.85, 11.65, "הזמנים טיפוסיים ל-20°C וצמנט רגיל; חום ומאיצים מקצרים, קור ומעכבים מאריכים")
message(s, "שגיאה בשלב הפלסטי (מים, ריטוט מאוחר) אינה ניתנת לתיקון", y=5.55, h=0.65)

# ================================================================ Section 3
divider(3, "חוזק הבטון", ["הגדרה", "התפתחות", "גורמים", "אשפרה", "דרגות חוזק"])

# ---------------------------------------------------------------- 13. What is strength
s = content_slide("מהו חוזק?")
pic_fit(s, "cube_load", 0.9, 1.2, 4.2, 4.6)
box(s, 5.6, 1.3, 6.95, 1.5, fill=NAVY)
text(s, 5.85, 1.35, 6.45, 1.4, "יכולת הבטון לשאת עומסי לחיצה ללא כשל", size=26, bold=True, color=WHITE,
     align="c", anchor="m")
box(s, 5.6, 3.05, 3.3, 1.75, fill=TINT)
text(s, 5.7, 3.15, 3.1, 1.6, [{"t": "f = F / A", "size": 30, "bold": True, "color": BLUE, "align": "c"},
                              {"t": "כוח השבר חלקי שטח החתך", "size": 14, "align": "c"},
                              {"t": "MPa = N/mm²", "size": 13, "color": MUTED, "align": "c"}], anchor="m")
box(s, 9.25, 3.05, 3.3, 1.75, fill=TINT)
text(s, 9.35, 3.15, 3.1, 1.6, [{"t": "דוגמה", "size": 15, "bold": True, "color": NAVY, "align": "c"},
                               {"t": "קובייה 150 מ\"מ נשברה ב-675 kN", "size": 14, "align": "c"},
                               {"t": "675,000 / 22,500 = 30 MPa", "size": 15, "bold": True, "color": BLUE,
                                "align": "c"}], anchor="m")
text(s, 5.6, 5.05, 6.95, 0.8, "חוזק הלחיצה הוא המדד המרכזי לאיכות הבטון, אבל לא היחיד: עמידות וחדירות נבדקות בנפרד",
     size=15, color=MUTED)

# ---------------------------------------------------------------- 14. Strength development
s = content_slide("התפתחות חוזק לאורך זמן")
cd = CategoryChartData()
cd.categories = ["1", "3", "7", "28", "90"]
curves = [("CEM I 52.5R", (40, 65, 80, 100, 105), NAVY),
          ("CEM I 42.5", (25, 52, 70, 100, 110), BLUE),
          ("CEM II/A", (18, 42, 62, 100, 115), GRAY),
          ("CEM III/A", (10, 30, 50, 100, 122), LIME_D)]
for nm, vals, _ in curves:
    cd.add_series(nm, vals)
gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, I(0.9), I(1.15), I(8.0), I(4.85), cd)
ch = gf.chart
style_chart(ch)
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.BOTTOM
ch.legend.include_in_layout = False
ch.legend.font.size = Pt(12)
for ser, (_, _, col) in zip(ch.plots[0].series, curves):
    series_line(ser, col)
axis_fmt(ch.value_axis, rl("% מחוזק 28 יום"), fmt='0"%"')
ch.value_axis.maximum_scale = 130
ch.value_axis.minimum_scale = 0
ch.value_axis.major_unit = 20
axis_fmt(ch.category_axis, rl("גיל (ימים)"), gridlines=False)
caption(s, 0.9, 6.05, 8.0, "52.5R: חוזק מוקדם גבוה · CEM II: עם אבן גיר / אפר פחם · CEM III: עם סיגים. ערכים טיפוסיים להמחשה")
pts = [("28", "יום — גיל הייחוס", "דרגת החוזק (למשל ב-30) נבדקת בקוביות בגיל 28 יום"),
       ("7", "ימים — אינדיקציה מוקדמת", "כ-65%-75% מחוזק 28 בצמנט רגיל"),
       ("90", "יום — המשך עלייה", "צמנטים עם תוספים ממשיכים לצבור חוזק")]
for i, (n, t1, t2) in enumerate(pts):
    y = 1.3 + i * 1.5
    text(s, 9.2, y, 3.35, 0.55, [[(n + " ", {"size": 30, "bold": True, "color": BLUE}),
                                   (t1, {"size": 15, "bold": True, "color": NAVY})]])
    text(s, 9.2, y + 0.6, 3.35, 0.8, t2, size=14)

# ---------------------------------------------------------------- 15. Factors on strength
s = content_slide("גורמים המשפיעים על החוזק")
facs = [("FaTint", "יחס W/C", "הגורם הדומיננטי: W/C נמוך יותר = חוזק גבוה יותר (חוק אברמס)", NAVY),
        ("FaWarehouse", "צמנט", "סוג, דרגת חוזק (42.5 / 52.5) ותכולה בתערובת", BLUE),
        ("FaCubes", "אגרגטים", "חוזק הסלע, צורה, ניקיון, גרדציה וגודל מקסימלי", BLUE),
        ("FaShieldAlt", "אשפרה", "בטון ללא אשפרה עלול להגיע לכ-50%-60% בלבד מהחוזק הפוטנציאלי", NAVY)]
for i, (ic, t1, t2, col) in enumerate(facs):
    x = 6.85 if i % 2 == 0 else 0.9
    y = 1.25 if i < 2 else 3.4
    box(s, x, y, 5.7, 1.95, fill=TINT)
    icon(s, ic, x + 5.7 - 0.6, y + 0.6, 0.8, fill=col)
    text(s, x + 0.25, y + 0.3, 4.4, 0.5, t1, size=21, bold=True, color=NAVY)
    text(s, x + 0.25, y + 0.9, 4.4, 0.95, t2, size=15)
message(s, "שלושה מתוך ארבעה גורמים נקבעים במפעל — האשפרה נקבעת באתר", y=5.75, h=0.65)

# ---------------------------------------------------------------- 16. Curing (added)
s = content_slide("אשפרה: הגורם שנשכח")
pic_fit(s, "curing", 0.9, 1.15, 4.6, 2.8)
cd = CategoryChartData()
cur = [("אשפרה רציפה", 100), ("שבוע אשפרה", 90), ("אשפרה קצרה", 80), ("ללא אשפרה", 60)]
cd.categories = [rl(c[0]) for c in reversed(cur)]
cd.add_series("חוזק יחסי", [c[1] for c in reversed(cur)])
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, I(5.75), I(1.15), I(6.8), I(2.9), cd)
ch = gf.chart
style_chart(ch, 12)
ch.has_legend = False
ch.has_title = True
ch.chart_title.text_frame.text = rl("חוזק יחסי בגיל 28 יום (%)")
for p in ch.chart_title.text_frame.paragraphs:
    for r in p.runs:
        set_font(r, 14, NAVY, True)
pl = ch.plots[0]
pl.gap_width = 60
pl.vary_by_categories = False
pl.series[0].format.fill.solid()
pl.series[0].format.fill.fore_color.rgb = rgb(BLUE)
pt = pl.series[0].points[0]
pt.format.fill.solid()
pt.format.fill.fore_color.rgb = rgb(GRAY)
pl.has_data_labels = True
pl.data_labels.number_format = '0"%"'
pl.data_labels.number_format_is_linked = False
pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
pl.data_labels.font.size = Pt(12)
pl.data_labels.font.bold = True
axis_fmt(ch.value_axis, gridlines=False)
ch.value_axis.visible = False
ch.value_axis.maximum_scale = 115
ch.value_axis.minimum_scale = 0
axis_fmt(ch.category_axis, gridlines=False)
caption(s, 5.75, 4.05, 6.8, "אשפרה קצרה = 3 ימים. מגמה טיפוסית לפי ספרות מקצועית — להמחשה")
meths = [("FaWater", "הרטבה רציפה", "המטרה, הצפה, יוטה רטובה"),
         ("FaLayerGroup", "כיסוי", "יריעות פוליאתילן צמודות"),
         ("FaVial", "חומר אשפרה", "Curing Compound מיד אחרי ההחלקה"),
         ("FaClock", "משך", "לפחות 7 ימים; יותר עם סיגים/אפר פחם או בקור")]
for i, (ic, t1, t2) in enumerate(meths):
    x = 12.55 - (i + 1) * 2.93 + 0.08
    box(s, x, 4.5, 2.8, 1.25, fill=TINT)
    icon(s, ic, x + 2.8 - 0.38, 4.88, 0.5)
    text(s, x + 0.15, 4.62, 2.0, 0.45, t1, size=16, bold=True, color=NAVY)
    text(s, x + 0.15, 5.1, 2.5, 0.62, t2, size=13)
text(s, 0.9, 5.9, 8.9, 0.6, [[("האשפרה משפיעה בעיקר על שכבת הכיסוי — השכבה שמגנה על הזיון. ", {"bold": True, "color": NAVY}),
                             ("משך ושיטה: לפי ת\"י 467 ומפרט הפרויקט", {"color": MUTED, "size": 13})]], size=15)

# ---------------------------------------------------------------- 17. Strength classes (ת"י 118)
s = content_slide("דרגות חוזק לפי ת\"י 118")
text(s, 0.9, 1.1, 11.65, 0.5, "דרגת החוזק מסומנת באות ב' ובמספר: המספר הוא החוזק האופייני (MPa) על קוביות בגיל 28 יום",
     size=16, color=MUTED)
rows = [["דרגה", "חוזק אופייני על קוביות (MPa)", "שימושים טיפוסיים"],
        ["ב-20", "20", "בטון רזה, מילוי, אלמנטים לא נושאים"],
        ["ב-30", "30", "יסודות, רצפות, מבני מגורים"],
        ["ב-40", "40", "תקרות, עמודים, סביבה מליחה"],
        ["ב-50", "50", "טרום, גשרים, סביבה ימית"],
        ["ב-60", "60", "בטון בחוזק גבוה, רבי-קומות"]]
table(s, 0.9, 1.75, 11.65, rows, [2.0, 3.6, 6.05], size=15, row_h=0.55, header_h=0.55,
      align=["c", "c", "r"], cell_fill=lambda i, j: LIME_T if j == 0 else None)
caption(s, 0.9, 5.15, 6.0, "דרגות ביניים (ב-25, ב-35, ב-45...) קיימות לפי התקן")
box(s, 0.9, 5.45, 8.95, 1.05, fill=TINT)
icon(s, "FaLightbulb", 9.35, 5.97, 0.55)
text(s, 1.1, 5.5, 7.85, 0.95, [{"t": "דרגת החוזק אינה בחירה חופשית:", "bold": True, "color": NAVY},
                               {"t": "ת\"י 118 קובע דרגת חוזק מינימלית לכל דרגת חשיפה — והמתכנן קובע את הדרגה בפועל."}],
     size=14, anchor="m")

# ---------------------------------------------------------------- 17b. Exposure classes (ת"י 118)
s = content_slide("דרגות חשיפה לפי ת\"י 118")
text(s, 0.9, 1.1, 11.65, 0.5, "ת\"י 118 מגדיר 11 דרגות חשיפה לפי סביבת הרכיב. המזמין חייב לציין את דרגת החשיפה של כל רכיב",
     size=16, color=MUTED)
zones = [("1", "פנים המבנה", "רכיבים מוגנים, יבשים", "DCE6F1"),
         ("2", "חוץ, 2 ק\"מ ומעלה מהים", "מעל 2 מ' מהקרקע", "A9C6E3"),
         ("3", "חוץ, במגע עם קרקע", "קרקע לא אגרסיבית, עד 2 מ' מעליה", "7FA7D6"),
         ("4", "חוץ, 1-2 ק\"מ מהים", "סביבה מליחה", BLUE),
         ("5", "חוץ, 200 מ' עד 1 ק\"מ מהים", "סביבה ימית", "1F5A94"),
         ("6-11", "ימי וכימי חמורים", "קו החוף, מים, קרקע וסביבה אגרסיבית", NAVY)]
zw = 11.65 / 6
for k, (n, t1, t2, col) in enumerate(zones):
    x = 12.55 - (k + 1) * zw
    r = box(s, x + 0.02, 1.75, zw - 0.04, 0.8, fill=col, shape=MSO_SHAPE.RECTANGLE)
    shape_text(r, [[("דרגה " + n, {"size": 18, "color": NAVY if k < 3 else WHITE})]])
    text(s, x + 0.08, 2.65, zw - 0.16, 0.45, t1, size=14, bold=True, color=NAVY, align="c")
    text(s, x + 0.08, 3.1, zw - 0.16, 0.65, t2, size=12, align="c")
line(s, 12.5, 3.85, 0.95, 3.85, color=MUTED, w=2)
text(s, 0.9, 3.92, 5.0, 0.3, "קרבה לים ואגרסיביות הסביבה ←", size=12, color=MUTED, align="l")
verify(s, 10.8, 3.95)
text(s, 0.9, 4.4, 11.65, 0.4, "מה נגזר מדרגת החשיפה", size=17, bold=True, color=NAVY)
derived = [("FaTint", "W/C מרבי"), ("FaWarehouse", "תכולת צמנט מינימלית"), ("FaWeightHanging", "דרגת חוזק מינימלית"),
           ("FaShieldAlt", "עובי כיסוי לזיון"), ("FaFlask", "תקרת כלורידים")]
for k, (ic, t1) in enumerate(derived):
    x = 12.55 - (k + 1) * 2.33 + 0.05
    box(s, x, 4.85, 2.23, 0.75, fill=TINT)
    icon(s, ic, x + 1.93, 5.22, 0.45, fill=BLUE)
    text(s, x + 0.1, 4.85, 1.5, 0.75, t1, size=13, bold=True, color=NAVY, anchor="m")
caption(s, 0.9, 5.75, 9.0, "תיאורי הדרגות מקוצרים. ההגדרה המלאה והערכים לכל דרגה — בטבלאות ת\"י 118 במהדורה המחייבת")

# ---------------------------------------------------------------- 18. Triangle
s = content_slide("משמעות החוזק: האיזון בתכן התערובת")
tri = box(s, 1.6, 1.4, 5.2, 4.3, fill=TINT, shape=MSO_SHAPE.ISOSCELES_TRIANGLE)
text(s, 2.7, 3.75, 3.0, 0.9, [{"t": "תכן תערובת", "size": 18, "bold": True, "color": NAVY},
                               {"t": "מאוזן", "size": 18, "bold": True, "color": NAVY}], align="c")
for ic, t1, cx_, cy_ in (("FaWeightHanging", "חוזק", 4.2, 1.35), ("FaShieldAlt", "עמידות", 6.8, 5.65),
                         ("FaHandPaper", "עבידות", 1.6, 5.65)):
    icon(s, ic, cx_, cy_, 1.0, fill=NAVY, ring=WHITE)
    if cy_ > 3:
        text(s, cx_ - 1.0, cy_ + 0.55, 2.0, 0.4, t1, size=18, bold=True, color=BLUE, align="c")
    else:
        text(s, cx_ + 0.6, cy_ - 0.2, 1.6, 0.4, t1, size=18, bold=True, color=BLUE, align="r")
rows_ = [("FaTint", "מוסיפים מים", "עבידות ↑  חוזק ↓  עמידות ↓", RED_T, RED),
         ("FaWarehouse", "מוסיפים צמנט בלבד", "חוזק ↑  עלות ↑  חום וסדיקה ↑", AMBER_T, AMBER),
         ("FaVial", "W/C נמוך + סופר-פלסטיצייזר", "חוזק ↑  עמידות ↑  עבידות נשמרת", GREEN_T, GREEN)]
for i, (ic, t1, t2, bg, col) in enumerate(rows_):
    y = 1.4 + i * 1.45
    box(s, 7.9, y, 4.65, 1.25, fill=bg)
    glyph(s, ic, 11.95, y + 0.35, 0.5, "b")
    text(s, 8.1, y + 0.12, 3.7, 0.45, t1, size=17, bold=True, color=NAVY)
    text(s, 8.1, y + 0.62, 3.7, 0.5, t2, size=15, bold=True, color=col)

# ================================================================ Section 4
divider(4, "יחס מים-צמנט", ["הגדרה", "W/C וחוזק", "W/C וחדירות", "W/C ועמידות", "מקרה שטח"])

# ---------------------------------------------------------------- 19. W/C definition
s = content_slide("יחס מים-צמנט (W/C)")
box(s, 3.4, 1.25, 9.15, 1.7, fill=NAVY)
text(s, 3.6, 1.3, 8.75, 1.6, [[("W/C  =  ", {"size": 34, "bold": True, "color": LIME}),
                               ("משקל המים (ק\"ג)  ÷  משקל הצמנט (ק\"ג)", {"size": 26, "bold": True, "color": WHITE})]],
     align="c", anchor="m")
icon(s, "FaBalanceScale", 2.1, 2.1, 1.5, fill=BLUE)
ex = [("180", "ליטר מים ל-מ\"ק"), ("÷", ""), ("360", "ק\"ג צמנט ל-מ\"ק"), ("=", ""), ("0.50", "W/C")]
xs = 11.0
for v, t_ in ex:
    w = 2.2 if t_ else 0.6
    xs -= w
    if t_:
        box(s, xs, 3.3, w, 1.35, fill=TINT)
        text(s, xs, 3.35, w, 1.25, [{"t": v, "size": 32, "bold": True, "color": BLUE if v != "0.50" else NAVY},
                                    {"t": t_, "size": 13}], align="c", anchor="m")
    else:
        text(s, xs, 3.3, w, 1.35, v, size=32, bold=True, color=MUTED, align="c", anchor="m")
    xs -= 0.15
notes = [("FaTint", "סופרים את כל המים", "כולל הלחות החופשית באגרגטים ומים מהמוספים"),
         ("FaRecycle", "W/B — מים לחומר מלכד", "כשיש תוספים: מחלקים בצמנט + אפר/סיגים/מיקרוסיליקה"),
         ("FaBook", "מגבלה תקנית", "ת\"י 118 קובע W/C מרבי לכל אחת מ-11 דרגות החשיפה")]
for i, (ic, t1, t2) in enumerate(notes):
    x = 12.55 - (i + 1) * 3.88 + 0.08
    icon(s, ic, x + 3.4, 5.25, 0.55, fill=NAVY)
    text(s, x, 4.95, 3.0, 0.4, t1, size=16, bold=True, color=NAVY)
    text(s, x, 5.35, 3.0, 0.95, t2, size=13)

# ---------------------------------------------------------------- 20. W/C vs strength
s = content_slide("W/C וחוזק")
cd = XyChartData()
ser = cd.add_series(rl("חוזק קובייה 28 יום"))
for wc, f in ((0.38, 55), (0.42, 50), (0.47, 44), (0.54, 37), (0.61, 31), (0.69, 25), (0.79, 19)):
    ser.add_data_point(wc, f)
gf = s.shapes.add_chart(XL_CHART_TYPE.XY_SCATTER_SMOOTH, I(0.9), I(1.15), I(7.6), I(4.8), cd)
ch = gf.chart
style_chart(ch)
ch.has_legend = False
series_line(ch.plots[0].series[0], BLUE, w=3.5, smooth=True)
axis_fmt(ch.value_axis, rl("חוזק קובייה (MPa)"))
ch.value_axis.minimum_scale, ch.value_axis.maximum_scale, ch.value_axis.major_unit = 10, 60, 10
axis_fmt(ch.category_axis, rl("יחס מים-צמנט"), gridlines=False, fmt="0.00")
ch.category_axis.minimum_scale, ch.category_axis.maximum_scale, ch.category_axis.major_unit = 0.35, 0.80, 0.05
caption(s, 0.9, 6.0, 7.6, "עקומה טיפוסית להמחשה — לכל מפעל ולכל צמנט עקומה משלו, שנבנית מתערובות ניסיון")
box(s, 8.85, 1.25, 3.7, 1.9, fill=NAVY)
text(s, 9.0, 1.3, 3.4, 1.8, [{"t": "ככל ש-W/C נמוך יותר", "size": 17, "color": WHITE, "bold": True},
                             {"t": "החוזק גבוה יותר", "size": 24, "color": LIME, "bold": True}], align="c",
     anchor="m")
box(s, 8.85, 3.35, 3.7, 1.25, fill=TINT)
text(s, 9.0, 3.4, 3.4, 1.15, [{"t": "W/C 0.47 → 0.61", "size": 22, "bold": True, "color": BLUE, "align": "c", "rtl": False},
                             {"t": "החוזק יורד בכ-30%, מ-44 ל-31 MPa", "size": 14, "align": "c"}], anchor="m")
text(s, 8.85, 4.8, 3.7, 1.0, "חוק אברמס (1918): החוזק תלוי ביחס המים-צמנט — כמעט בלי קשר לכמות הצמנט עצמה",
     size=13, color=MUTED)

# ---------------------------------------------------------------- 21. W/C and permeability
s = content_slide("W/C וחדירות")
pic_fit(s, "pores_dense", 6.95, 1.6, 5.6, 3.3)
pic_fit(s, "pores_open", 0.9, 1.6, 5.6, 3.3)
text(s, 6.95, 1.15, 5.6, 0.4, [[("W/C 0.40", {"bold": True, "color": NAVY, "size": 19}),
                                  ("  —  משחה צפופה", {"size": 17})]], align="c")
text(s, 0.9, 1.15, 5.6, 0.4, [[("W/C 0.70", {"bold": True, "color": RED, "size": 19}),
                                 ("  —  נקבוביות רציפות", {"size": 17})]], align="c")
text(s, 6.95, 5.0, 5.6, 0.75, [{"t": "נקבוביות קפילריות מעטות ומבודדות", "bullet": True},
                               {"t": "מים ויונים מתקשים לחדור", "bullet": True}], size=15)
text(s, 0.9, 5.0, 5.6, 0.75, [{"t": "רשת תעלות מחוברות לכל העומק", "bullet": True},
                              {"t": "מים, כלורידים ו-CO₂ חודרים בקלות", "bullet": True}], size=15)
message(s, "מעל W/C ≈ 0.6 הנקבוביות הופכות לרציפות — והחדירות עולה בסדרי גודל", y=5.85, h=0.62)

# ---------------------------------------------------------------- 22. W/C and durability
s = content_slide("W/C ועמידות")
pic_fit(s, "chloride", 0.9, 1.2, 7.3, 3.2)
text(s, 0.9, 4.45, 7.3, 0.6, "קו מקווקו = עומק כיסוי הבטון. משמאל: הכלורידים הגיעו לזיון, החלודה מתנפחת ומבקיעה את הבטון",
     size=13, color=MUTED)
imp = [("FaFlask", "כלורידים", "חודרים דרך הנקבוביות ומשמידים את שכבת ההגנה של הפלדה"),
       ("FaBolt", "קורוזיה", "חלודה תופסת פי 2-6 מנפח הפלדה — סדיקה והתפוררות"),
       ("FaTint", "חדירת מים", "מובילה מלחים, סולפטים ותגובות אלקלי-סיליקה")]
for i, (ic, t1, t2) in enumerate(imp):
    y = 1.25 + i * 1.35
    box(s, 8.5, y, 4.05, 1.2, fill=TINT)
    icon(s, ic, 12.05, y + 0.6, 0.62, fill=NAVY)
    text(s, 8.65, y + 0.1, 3.0, 0.4, t1, size=17, bold=True, color=NAVY)
    text(s, 8.65, y + 0.5, 3.0, 0.68, t2, size=13)
message(s, "בדרגות החשיפה הימיות של ת\"י 118: W/C נמוך וכיסוי מוגדל — זו ההגנה העיקרית לזיון", y=5.25, h=0.85)

# ---------------------------------------------------------------- 23. Case study
s = content_slide("מקרה שטח: הוספת מים באתר")
cols3 = [("FaExclamationTriangle", "הבעיה", RED, RED_T,
          ["המשאית הגיעה לאחר המתנה ארוכה, בשקיעה נמוכה", "הצוות ביקש \"להוסיף קצת מים\" כדי לשאוב"]),
         ("FaSearch", "הגורם", AMBER, AMBER_T,
          ["תכנון לוגיסטי לקוי וזמן המתנה", "אין נוהל לתיקון עבידות במוסף", "לחץ זמנים — ללא תיעוד"]),
         ("FaArrowDown", "התוצאה", NAVY, TINT,
          ["חוזק נמוך מהדרישה בקוביות 28 יום", "חדירות גבוהה ואבק פני שטח", "מחלוקת אחריות מול הלקוח"])]
for i, (ic, t1, col, bg, items) in enumerate(cols3):
    x = 12.55 - (i + 1) * 3.95 + 0.1
    box(s, x, 1.2, 3.85, 2.75, fill=bg)
    icon(s, ic, x + 3.3, 1.7, 0.7, fill=col)
    text(s, x + 0.2, 1.48, 2.6, 0.45, t1, size=21, bold=True, color=col)
    text(s, x + 0.2, 2.2, 3.45, 1.7, [{"t": it, "bullet": True} for it in items], size=14)
    if i < 2:
        line(s, x + 0.02, 2.6, x - 0.08, 2.6, color=GRAY, w=3)
box(s, 0.9, 4.15, 11.65, 1.3, fill=TINT2)
calc = [("300 ק\"ג צמנט, 165 ל' מים", "W/C = 0.55", NAVY), ("+ 20 ל' מים למ\"ק", "W/C = 0.62", RED),
        ("ירידת חוזק צפויה (בקירוב)", "15%-20%", RED)]
for i, (t1, t2, col) in enumerate(calc):
    x = 12.55 - (i + 1) * 3.88 + 0.08
    text(s, x, 4.22, 3.7, 0.4, t1, size=14, color=MUTED, align="c")
    text(s, x, 4.65, 3.7, 0.7, t2, size=26, bold=True, color=col, align="c")
message(s, "הפתרון: מוסף מאושר בהנחיית המפעל ותיעוד — לעולם לא מים", y=5.65, h=0.7,
        ic="FaCheckCircle")

# ================================================================ Section 5
divider(5, "בדיקות בטון טרי", ["שקיעה", "טמפרטורה", "צפיפות", "כמות אוויר"])

# ---------------------------------------------------------------- 24. Fresh concrete
s = content_slide("מהו בטון טרי?")
text(s, 0.9, 1.15, 11.65, 0.5, "בטון במצב פלסטי, מרגע הערבוב ועד תחילת ההתקשרות. הבדיקות כאן הן ההזדמנות האחרונה לעצור משאית",
     size=16, color=MUTED)
tests = [("FaRulerVertical", "שקיעה", "עבידות ועקביות בין משאיות", "ת\"י 26 חלק 2"),
         ("FaThermometerHalf", "טמפרטורה", "סיכון להתקשרות מהירה או איטית", "בקרת ייצור לפי ת\"י 118"),
         ("FaWeightHanging", "צפיפות", "בקרת תפוקה ופרופורציות", "ת\"י 26 חלק 2"),
         ("FaTachometerAlt", "כמות אוויר", "עמידות, חוזק ותפוקה", "ת\"י 26 חלק 2")]
for i, (ic, t1, t2, st) in enumerate(tests):
    x = 12.55 - (i + 1) * 2.92 + 0.07
    box(s, x, 1.85, 2.78, 3.35, fill=TINT)
    icon(s, ic, x + 1.39, 2.7, 1.1, fill=BLUE if i % 2 == 0 else NAVY)
    text(s, x + 0.1, 3.45, 2.58, 0.45, t1, size=21, bold=True, color=NAVY, align="c")
    text(s, x + 0.15, 3.95, 2.48, 0.7, t2, size=14, align="c")
    text(s, x + 0.1, 4.7, 2.58, 0.4, st, size=12, color=MUTED, align="c")
message(s, "בודקים את המשאית הראשונה, כל משאית חשודה — ולפי תכנית הדגימה", y=5.55, h=0.7)

# ---------------------------------------------------------------- 25. Slump test
s = content_slide("בדיקת שקיעה (Slump Test)")
pic_fit(s, "slump_cone", 0.9, 1.2, 6.2, 3.6)
text(s, 0.9, 4.85, 6.2, 0.35, "מידות החרוט במ\"מ: קוטר עליון 100, תחתון 200, גובה 300", size=13, color=MUTED,
     align="c")
steps = ["מרטיבים חרוט ומשטח, מניחים על בסיס יציב ואופקי",
         "ממלאים ב-3 שכבות שוות, כל אחת 25 דקירות במוט",
         "מיישרים את פני הבטון ומנקים סביב החרוט",
         "מרימים אנכית ובאחידות, ללא סיבוב (2-5 שניות)",
         "מודדים מגובה החרוט עד מרכז הבטון השוקע",
         "רושמים את התוצאה ואת צורת השקיעה"]
for i, st in enumerate(steps):
    y = 1.25 + i * 0.72
    badge(s, i + 1, 12.25, y + 0.28, 0.48, fill=BLUE)
    text(s, 7.5, y, 4.4, 0.62, st, size=15, anchor="m")
message(s, "הבדיקה לפי ת\"י 26 חלק 2 — מבוצעת ברצף ובמהירות", x=0.9, w=6.2, y=5.4, h=0.7, ic="FaClock")

# ---------------------------------------------------------------- 26. Slump interpretation
s = content_slide("פענוח תוצאות שקיעה")
kinds = [("slump_low", "נמוכה מדי", AMBER, AMBER_T, "קשה לשאיבה ולריטוט; סיכון לקיני חצץ",
          "בדיקת זמן המתנה; תיקון רק במוסף מאושר"),
         ("slump_ok", "תקינה", GREEN, GREEN_T, "שקיעה סימטרית בטווח ההזמנה",
          "מאשרים ומתעדים"),
         ("slump_high", "גבוהה מדי", RED, RED_T, "חשד לעודף מים; סגרגציה ודימום",
          "עצירה, בדיקה חוזרת, החלטה על דחייה")]
for i, (img, t1, col, bg, t2, t3) in enumerate(kinds):
    x = 12.55 - (i + 1) * 3.88 + 0.08
    box(s, x, 1.2, 3.8, 4.1, fill=bg)
    pic_fit(s, img, x + 0.3, 1.35, 3.2, 1.75)
    text(s, x + 0.15, 3.2, 3.5, 0.45, t1, size=22, bold=True, color=col, align="c")
    text(s, x + 0.2, 3.7, 3.4, 0.7, t2, size=14, align="c")
    text(s, x + 0.2, 4.45, 3.4, 0.75, [[("פעולה: ", {"bold": True, "color": NAVY}), (t3, {})]], size=13,
         align="c")
rows = [["S1", "S2", "S3", "S4", "S5"], ["10-40", "50-90", "100-150", "160-210", "220 ומעלה"]]
text(s, 9.6, 5.45, 2.95, 0.35, "דרגות סומך לפי ת\"י 118 (מ\"מ)", size=13, bold=True, color=NAVY)
verify(s, 10.8, 5.85)
table(s, 0.9, 5.45, 8.5, rows, [1, 1, 1, 1, 1], size=13, row_h=0.4, header_h=0.4, align=["c"] * 5,
      bold_first=False)

# ---------------------------------------------------------------- 27. Temperature
s = content_slide("בדיקת טמפרטורה")
zones = [(0, 5, "4A6FA5", "סכנת קפיאה", "ההידרציה כמעט נעצרת"),
         (5, 10, "7FA7D6", "איטי", "התקשרות ממושכת"),
         (10, 30, GREEN, "תקין", "טווח עבודה רגיל"),
         (30, 35, AMBER, "חם", "נדרשים אמצעי קיץ"),
         (35, 40, RED, "סיכון גבוה", "בדרך כלל מחוץ למפרט")]
sx, sw = 0.9, 11.65
for a, b_, col, t1, t2 in zones:
    x = sx + sw - (b_ / 40) * sw
    w = (b_ - a) / 40 * sw
    r = box(s, x, 1.55, w - 0.03, 0.9, fill=col, shape=MSO_SHAPE.RECTANGLE)
    shape_text(r, t1, size=15 if w > 1.2 else 13)
    text(s, x, 2.52, w - 0.03, 0.55, t2, size=12, color=MUTED, align="c")
for v in (0, 5, 10, 30, 35, 40):
    x = sx + sw - (v / 40) * sw
    text(s, x - 0.35, 1.18, 0.7, 0.32, f"{v}°C", size=13, bold=True, color=NAVY, align="c", rtl=False)
eff = [("FaSun", "בטון חם", ["התקשרות מהירה ואיבוד שקיעה", "סדיקה פלסטית ותרמית", "חוזק סופי נמוך יותר"], RED),
       ("FaSnowflake", "בטון קר", ["התקשרות והתפתחות חוזק איטיות", "פירוק תבניות מאוחר", "סכנת קפיאה בגיל צעיר"], BLUE),
       ("FaThermometerHalf", "איך מודדים", ["מדחום מכויל, טבילה של 75 מ\"מ לפחות", "2-5 דקות עד התייצבות", "רישום גם של טמפרטורת האוויר"], NAVY)]
for i, (ic, t1, items, col) in enumerate(eff):
    x = 12.55 - (i + 1) * 3.88 + 0.08
    box(s, x, 3.3, 3.8, 2.3, fill=TINT)
    icon(s, ic, x + 3.3, 3.8, 0.62, fill=col)
    text(s, x + 0.2, 3.55, 2.7, 0.45, t1, size=18, bold=True, color=NAVY)
    text(s, x + 0.2, 4.2, 3.4, 1.4, [{"t": it, "bullet": True} for it in items], size=13)
caption(s, 0.9, 5.75, 9.0, "הגבולות בפועל נקבעים במפרט ובתקן החל על הפרויקט; הטווחים כאן להמחשה")

# ---------------------------------------------------------------- 28. Density
s = content_slide("בדיקת צפיפות")
pic_fit(s, "density", 0.9, 1.25, 3.6, 3.8)
box(s, 4.85, 1.25, 7.7, 1.9, fill=TINT)
text(s, 5.05, 1.3, 7.3, 0.45, "מה מטרתה?", size=19, bold=True, color=NAVY)
text(s, 5.05, 1.8, 7.3, 1.3, [{"t": "בדיקת תפוקה — האם הלקוח מקבל מ\"ק מלא", "bullet": True},
                              {"t": "אימות פרופורציות התערובת ואינדיקציה לתכולת אוויר", "bullet": True},
                              {"t": "טיפוסי: 2,300-2,450 ק\"ג/מ\"ק לבטון רגיל", "bullet": True}], size=14)
text(s, 4.85, 3.35, 7.7, 0.45, "איך מבצעים?", size=19, bold=True, color=NAVY)
how = ["שוקלים מיכל ריק בנפח ידוע", "ממלאים בשכבות ומהדקים", "מיישרים ומנקים את הדפנות", "שוקלים ומחשבים"]
for i, t_ in enumerate(how):
    x = 12.55 - (i + 1) * 1.93 + 0.05
    badge(s, i + 1, x + 0.94, 4.1, 0.5, fill=BLUE)
    text(s, x, 4.45, 1.88, 0.75, t_, size=13, align="c")
box(s, 4.85, 5.3, 7.7, 0.85, fill=NAVY)
text(s, 5.0, 5.33, 7.4, 0.8, [{"t": "ρ = (M₂ − M₁) / V", "bold": True, "color": WHITE, "size": 20, "rtl": False},
                              {"t": "M₂ מיכל מלא  •  M₁ מיכל ריק  •  V נפח המיכל", "color": WHITE, "size": 12}],
     align="c", anchor="m", space_after=0)

# ---------------------------------------------------------------- 29. Air content
s = content_slide("כמות אוויר")
pic_fit(s, "air_meter", 0.9, 1.25, 3.5, 4.0)
caption(s, 0.9, 5.3, 3.5, "מד אוויר בשיטת הלחץ", align="c")
box(s, 8.75, 1.25, 3.8, 3.4, fill=GREEN_T)
icon(s, "FaCheckCircle", 12.1, 1.7, 0.6, fill=GREEN)
text(s, 8.95, 1.5, 2.7, 0.45, "יתרונות", size=20, bold=True, color=GREEN)
text(s, 8.95, 2.2, 3.4, 2.4, [{"t": "עמידות למחזורי קפיאה-הפשרה", "bullet": True},
                              {"t": "שיפור עבידות וקוהזיה", "bullet": True},
                              {"t": "הפחתת דימום וסגרגציה", "bullet": True}], size=14)
box(s, 4.75, 1.25, 3.8, 3.4, fill=RED_T)
icon(s, "FaTimesCircle", 8.1, 1.7, 0.6, fill=RED)
text(s, 4.95, 1.5, 2.7, 0.45, "חסרונות", size=20, bold=True, color=RED)
text(s, 4.95, 2.2, 3.4, 2.4, [{"t": "כל 1% אוויר ≈ 5% ירידה בחוזק", "bullet": True},
                              {"t": "השפעה על התפוקה והצפיפות", "bullet": True},
                              {"t": "רגישות לסוג המוסף ולזמן הערבוב", "bullet": True}], size=14)
box(s, 4.75, 4.85, 7.8, 1.3, fill=TINT)
text(s, 4.95, 4.9, 7.4, 1.2, [{"t": "בבטון רגיל בישראל: אוויר כלוא של 1%-2%. אוויר מוחדר מכוון — רק כשנדרש במפרט.",
                               "bold": True, "color": NAVY},
                              {"t": "תוצאה גבוהה בלי מוסף מחדיר אוויר = סימן לבעיה בתערובת"}], size=14, anchor="m")

# ================================================================ Section 6
divider(6, "בדיקות בטון קשוח", ["קוביות, קרוטים ובדיקות לא הורסות", "בדיקת לחיצה", "פענוח תוצאות"])

# ---------------------------------------------------------------- 30. Hardened concrete
s = content_slide("בדיקות בטון קשוח")
specs = [("spec_cube", "קוביות", "נלקחות מהבטון הטרי", "בקרת ההתאמה השוטפת של הבטון שסופק לדרגת החוזק", "ת\"י 26 חלקים 1 ו-4"),
         ("spec_core", "קרוטים", "קידוח מהמבנה", "הערכת הבטון בפועל במבנה, במקרי ספק או כשל", "ת\"י 26 חלק 6"),
         ("rebound", "בדיקות לא הורסות", "פטיש רתע ושיטות נוספות", "סקירה מהירה של אחידות; לא תחליף לקרוטים", "ת\"י 26 חלק 7")]
for i, (img, t1, t2, t3, st) in enumerate(specs):
    x = 12.55 - (i + 1) * 3.88 + 0.08
    box(s, x, 1.2, 3.8, 4.95, fill=TINT)
    pic_fit(s, img, x + 0.6, 1.35, 2.6, 2.3)
    text(s, x + 0.15, 3.75, 3.5, 0.45, t1, size=22, bold=True, color=NAVY, align="c")
    text(s, x + 0.15, 4.2, 3.5, 0.35, t2, size=15, bold=True, color=BLUE, align="c")
    text(s, x + 0.25, 4.62, 3.3, 0.9, t3, size=13, align="c")
    text(s, x + 0.15, 5.65, 3.5, 0.35, st, size=12, color=MUTED, align="c")

# ---------------------------------------------------------------- 31. Compression test
s = content_slide("בדיקת לחיצה")
pic_fit(s, "press", 0.9, 1.15, 4.0, 4.6)
steps = [("דגימה", "מהמשאית לפי תכנית, ת\"י 26 חלק 1"),
         ("יציקה והידוק", "בתבניות נקיות ומשומנות"),
         ("24 שעות ראשונות", "באתר, מוגן מזעזועים, חום והתייבשות"),
         ("אשפרה במעבדה", "במים או בתא לחות, בטמפרטורה מבוקרת"),
         ("לחיצה בגיל הבדיקה", "בקצב העמסה קבוע עד לשבר, ת\"י 26 חלק 4"),
         ("רישום ופענוח", "עומס, חוזק, צורת השבר וחריגות")]
for i, (t1, t2) in enumerate(steps):
    col = i % 2
    row = i // 2
    x = 8.85 - col * 3.9
    y = 1.25 + row * 1.55
    box(s, x, y, 3.7, 1.35, fill=TINT)
    badge(s, i + 1, x + 3.3, y + 0.38, 0.5, fill=BLUE)
    text(s, x + 0.2, y + 0.15, 2.8, 0.45, t1, size=17, bold=True, color=NAVY)
    text(s, x + 0.2, y + 0.62, 3.3, 0.7, t2, size=13)
text(s, 5.15, 5.95, 7.4, 0.5, "שבר תקין בקובייה: שתי פירמידות הפוכות. שבר חריג מעיד בדרך כלל על הכנה או אשפרה לקויות",
     size=13, color=MUTED)

# ---------------------------------------------------------------- 32. Cubes vs cores
s = content_slide("קוביות מול קרוטים")
rows = [["פרמטר", "קוביות", "קרוטים"],
        ["מה נבדק", "הבטון כפי שסופק", "הבטון כפי שנמצא במבנה"],
        ["תקן", "ת\"י 26 חלקים 1 ו-4", "ת\"י 26 חלק 6"],
        ["מתי", "שוטף, לפי תכנית הדגימה", "בספק, בכשל קוביות או בבדיקת מבנה קיים"],
        ["משקף", "את פוטנציאל התערובת באשפרה מבוקרת", "גם את הביצוע: הידוק, אשפרה ופגמים"],
        ["גורמים משפיעים", "דגימה, הכנה, אשפרה ראשונית", "כיוון וקוטר הקידוח, זיון, לחות"],
        ["פענוח", "מול דרגת החוזק לפי ת\"י 118", "לפי הקריטריונים לקרוטים בתקן"]]
table(s, 0.9, 1.25, 11.65, rows, [2.6, 4.5, 4.55], size=15, row_h=0.58, header_h=0.58,
      align=["r", "c", "c"], cell_fill=lambda i, j: TINT if j == 0 else None)
message(s, "קרוט נמוך מקובייה אינו בהכרח בטון פסול: הפער משקף לעיתים את הביצוע והאשפרה", y=5.75, h=0.65)

# ---------------------------------------------------------------- 33. Interpreting results
s = content_slide("פענוח תוצאות חוזק")
res = [38, 36, 40, 37, 39, 35, 38, 41, 37, 36, 38, 28, 37, 36, 35, 34, 35, 33, 34, 33]
mov = [None, None] + [round(sum(res[i - 2:i + 1]) / 3, 1) for i in range(2, len(res))]
cd = CategoryChartData()
cd.categories = [str(i + 1) for i in range(len(res))]
cd.add_series(rl("תוצאה בודדת"), res)
cd.add_series(rl("ממוצע נע (3)"), mov)
cd.add_series(rl("דרגה ב-30"), [30] * len(res))
cd.add_series(rl("יעד ייצור"), [37] * len(res))
gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, I(0.9), I(1.15), I(7.7), I(4.7), cd)
ch = gf.chart
style_chart(ch)
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.BOTTOM
ch.legend.include_in_layout = False
ser = ch.plots[0].series
series_line(ser[0], BLUE, w=1.75)
series_line(ser[1], NAVY, w=3, marker=False)
series_line(ser[2], RED, w=2, marker=False, dash=True)
series_line(ser[3], GREEN, w=1.5, marker=False, dash=True)
pt = ser[0].points[11]
pt.marker.format.fill.solid()
pt.marker.format.fill.fore_color.rgb = rgb(RED)
pt.marker.size = 12
axis_fmt(ch.value_axis, rl("חוזק קובייה (MPa)"))
ch.value_axis.minimum_scale, ch.value_axis.maximum_scale, ch.value_axis.major_unit = 24, 44, 4
axis_fmt(ch.category_axis, rl("מספר דגימה"), gridlines=False)
caption(s, 0.9, 5.9, 7.7, "דוגמה להמחשה: בטון ב-30. קריטריוני ההתאמה המחייבים — לפי ת\"י 118")
conc = [("FaChartLine", "Trend — מגמה", "ירידה עקבית בממוצע הנע (דגימות 13-20): לחקור חומרי גלם, מינון, ציוד"),
        ("FaExclamationTriangle", "Outlier — חריג", "דגימה 12 בודדת ונמוכה: לבדוק דגימה, הכנה ואשפרה לפני מסקנות"),
        ("FaBalanceScale", "בקרה סטטיסטית", "ממוצע, סטיית תקן ותוצאות בודדות מול קריטריוני ההתאמה של ת\"י 118")]
for i, (ic, t1, t2) in enumerate(conc):
    y = 1.2 + i * 1.6
    box(s, 8.85, y, 3.7, 1.45, fill=TINT)
    icon(s, ic, 12.1, y + 0.42, 0.55, fill=NAVY)
    text(s, 9.0, y + 0.15, 2.75, 0.45, t1, size=16, bold=True, color=NAVY)
    text(s, 9.0, y + 0.72, 3.35, 0.72, t2, size=12)

# ================================================================ Section 7
divider(7, "כשלים נפוצים", ["סגרגציה", "דימום", "קיני חצץ", "סדיקה פלסטית", "התקלפות"])


def defect(title, en, img, causes, prevent, note=None):
    s = content_slide(f"{title}  |  {en}")
    pic_fit(s, img, 0.9, 1.2, 5.9, 4.3)
    if note:
        caption(s, 0.9, 5.55, 5.9, note, align="c")
    for k, (hdr, items, col, bg, ic) in enumerate((("גורמים", causes, RED, RED_T, "FaSearch"),
                                                     ("מניעה", prevent, GREEN, GREEN_T, "FaShieldAlt"))):
        y = 1.2 if k == 0 else 3.75
        h = 2.4
        box(s, 7.1, y, 5.45, h, fill=bg)
        icon(s, ic, 12.05, y + 0.45, 0.55, fill=col)
        text(s, 7.3, y + 0.2, 4.2, 0.45, hdr, size=19, bold=True, color=col)
        text(s, 7.3, y + 0.8, 5.05, h - 0.9, [{"t": it, "bullet": True} for it in items], size=14, space_after=2)
    return s


defect("סגרגציה", "Segregation", "segregation",
       ["הטלה מגובה רב (מעל 1.5 מ') או דרך מזחלה ארוכה", "עודף מים / שקיעה גבוהה מדי", "ריטוט יתר או גרירת בטון בוויברטור",
        "גרדציה לקויה, מחסור בחומר דק"],
       ["שרוול או משפך עד קרוב לנקודת היציקה", "שקיעה בטווח ההזמנה, ללא תוספת מים", "ריטוט נקודתי וקצר",
        "תכן תערובת קוהזיבי"],
       "האגרגט הגס נפרד ומצטבר בתחתית; משחה עשירה נשארת למעלה")
defect("דימום", "Bleeding", "bleeding",
       ["W/C גבוה ומחסור בחומר דק", "התקשרות איטית (קור, מעכב במינון יתר)", "יציקות עמוקות", "אגרגט עם גרדציה פתוחה"],
       ["תכן עם מספיק דקים או תוספים", "מינון מוספים לפי הנחיה", "לא להחליק כשיש מים על הפנים", "יציקה בשכבות"],
       "מים עולים לפני השטח ונלכדים מתחת לאגרגטים ולמוטות הזיון")
defect("קיני חצץ", "Honeycombs", "honeycomb",
       ["ריטוט חסר או לא אחיד", "צפיפות זיון מול גודל אגרגט מרבי", "דליפת משחה מתבניות לא אטומות", "עבידות נמוכה וסגרגציה"],
       ["תכנית ריטוט: מרחקים ועומקים", "התאמת גודל האגרגט לצפיפות הזיון", "אטימת מישקי תבניות",
        "שקיעה מתאימה לאלמנט"],
       "אזור עם אגרגט חשוף וחללים, בדרך כלל בתחתית קירות ועמודים")
defect("סדיקה פלסטית", "Plastic Shrinkage", "plastic",
       ["רוח — מגבירה התאדות", "חום וקרינת שמש", "אידוי מהיר מקצב הדימום (מעל ~1 ק\"ג/מ\"ר/שעה)", "לחות יחסית נמוכה"],
       ["שוברי רוח וצל", "ערפול מים או חומר מעכב התאדות", "יציקה בשעות קרירות", "אשפרה מיד אחרי ההחלקה"],
       "סדקים קצרים ומקבילים בשעות הראשונות, לפני ההתקשות")
defect("התקלפות", "Scaling", "scaling",
       ["החלקה בזמן שיש מי דימום על הפנים", "החלקת יתר או פיזור צמנט יבש על הפנים", "אשפרה חסרה בשכבה העליונה",
        "W/C גבוה בפני השטח"],
       ["החלקה רק אחרי היעלמות מי הדימום", "אשפרה מיידית ומלאה", "W/C נמוך ובטון קוהזיבי",
        "תיקון: השחזה ושכבה מתקנת ייעודית"],
       "שכבה דקה עליונה מתקלפת וחושפת את האגרגט")

# ================================================================ Section 8
divider(8, "בקרת חומרי גלם", ["תהליך קבלה", "בדיקות אגרגטים", "צמנט", "מוספים", "זיהוי חומרים פסולים", "מטריצת סיכונים"])

# ---------------------------------------------------------------- 38. Raw-material control flow
s = content_slide("בקרת חומרי גלם: תהליך הקבלה")
flow = [("FaTruck", "קבלה", "תעודת משלוח, זיהוי ספק ותאריך, בדיקה ויזואלית"),
        ("FaFlask", "בדיקה", "דגימה ובדיקות לפי תכנית הבקרה"),
        ("FaClipboardCheck", "אישור", "השוואה לדרישות התקן והמפרט"),
        ("FaIndustry", "שימוש", "שחרור לממגורה או לערמה, עם עקיבות")]
for i, (ic, t1, t2) in enumerate(flow):
    x = 12.55 - (i + 1) * 2.95 + 0.1
    box(s, x, 1.3, 2.75, 2.3, fill=NAVY if i == 2 else TINT)
    icon(s, ic, x + 1.375, 1.9, 0.8, fill=LIME if i == 2 else BLUE, color="n" if i == 2 else "w")
    text(s, x + 0.1, 2.4, 2.55, 0.45, t1, size=20, bold=True, color=WHITE if i == 2 else NAVY, align="c")
    text(s, x + 0.15, 2.85, 2.45, 0.7, t2, size=13, color=WHITE if i == 2 else TEXT, align="c")
    if i < 3:
        line(s, x - 0.01, 2.45, x - 0.19, 2.45, color=BLUE, w=3)
# decision under "בדיקה": fail -> quarantine (right), pass -> up into "אישור"
bx_test = 12.55 - 2 * 2.95 + 0.1          # left edge of the "בדיקה" card
cx_test, cx_appr = bx_test + 1.375, bx_test - 2.95 + 1.375
dia = box(s, cx_test - 1.0, 3.9, 2.0, 1.15, fill=AMBER_T, line=AMBER, lw=2, shape=MSO_SHAPE.DIAMOND)
shape_text(dia, "תקין?", size=16, color=NAVY)
line(s, cx_test, 3.62, cx_test, 3.9, color=BLUE, w=2.5)
line(s, cx_test + 1.0, 4.47, 9.95, 4.47, color=RED, w=2.5)
text(s, cx_test + 1.0, 4.05, 0.7, 0.35, "לא", size=13, bold=True, color=RED, align="c")
r = box(s, 9.95, 3.9, 2.6, 1.15, fill=RED_T, line=RED, lw=1.5)
shape_text(r, [[("הסגר וסימון", {"bold": True, "color": RED, "size": 15})],
               [("החזרה לספק + דוח אי-התאמה", {"bold": False, "color": TEXT, "size": 12})]])
line(s, cx_test - 1.0, 4.47, cx_appr, 4.47, color=GREEN, w=2.5, head=False)
line(s, cx_appr, 4.47, cx_appr, 3.62, color=GREEN, w=2.5)
text(s, cx_test - 1.7, 4.05, 0.7, 0.35, "כן", size=13, bold=True, color=GREEN, align="c")
text(s, 0.9, 3.85, 4.1, 1.6, [{"t": "חומר שלא אושר — לא נכנס לממגורה או לערמה הפעילה", "bullet": True},
                              {"t": "כל משלוח ניתן לאיתור עד היציקה", "bullet": True},
                              {"t": "תיעוד: מי בדק, מתי ומה הוחלט", "bullet": True}], size=13)
message(s, "המקום הזול ביותר לעצור בעיה הוא בשער המפעל", y=5.65, h=0.7)

# ---------------------------------------------------------------- Aggregate tests (8 slides)
REQ_NOTE = "ערכים מקובלים — לאימות מול ת\"י 3 במהדורה המחייבת (כולל גיליונות התיקון) ומול מפרט המפעל"


def steps_list(s, items, x, y, w, gap=0.72, size=14, d=0.46, fill=BLUE):
    for i, it in enumerate(items):
        yy = y + i * gap
        badge(s, i + 1, x + w - d / 2, yy + 0.28, d, fill=fill, size=14)
        text(s, x, yy, w - d - 0.15, gap - 0.06, it, size=size, anchor="m")


# 1 — test map
s = content_slide("בדיקות אגרגטים: מפת הבדיקות")
rows = [["בדיקה", "דרישה לפי", "מה היא מגלה", "תדירות מומלצת"],
        ["דגימה והקטנה", "ת\"י 3", "הבסיס לכל תוצאה אמינה", "בכל בדיקה"],
        ["עקום דירוג ואבק", "ת\"י 3", "פיזור גדלים ותכולת אבק", "שבועי ובכל משלוח חשוד"],
        ["שווה חול / כחול מתילן", "ת\"י 3 / מפרט", "האם הדקים הם חרסית מזיקה", "שבועי (חול)"],
        ["לחות", "נוהל המפעל", "תיקון מים יומי בתערובת", "יומי ובכל שינוי"],
        ["צפיפות וספיגות", "ת\"י 3", "איכות הסלע, בסיס לתיקון מים", "חודשי"],
        ["שחיקה וצורת גרגר", "ת\"י 3", "עמידות, פחוסים ומוארכים", "רבעוני ובמקור חדש"],
        ["חומרים מזיקים", "ת\"י 3, ת\"י 118", "אורגני, חרסית, כלורידים, סולפטים", "חצי-שנתי ובמקור חדש"]]
table(s, 0.9, 1.2, 11.65, rows, [2.4, 3.0, 3.6, 2.65], size=14, row_h=0.53, header_h=0.5,
      align=["r", "c", "r", "r"], cell_fill=lambda i, j: TINT if j == 0 else None)
message(s, "התדירות נקבעת בתכנית הבקרה של המפעל — ומתהדקת בכל שינוי מקור או חריגה", y=5.75, h=0.65,
        ic="FaListOl")

# 2 — sampling
s = content_slide("דגימה והקטנת דגימה")
pic_fit(s, "sampling", 0.9, 1.25, 6.4, 3.0)
text(s, 0.9, 4.3, 3.0, 0.35, "נקודות דגימה בערמה", size=13, color=MUTED, align="c")
text(s, 4.2, 4.3, 3.1, 0.35, "מחלק (Riffle) לשני חצאים", size=13, color=MUTED, align="c")
steps_list(s, ["דוגמים מכמה נקודות בערמה, מתחת לשכבת הפנים — לא מהשוליים שבהם הגס מתגלגל",
               "מהמסוע: חתך מלא של הסרט, כשהמסוע עוצר",
               "מאחדים לדגימה מייצגת בכמות מספקת לגודל הגרגר המרבי",
               "מקטינים במחלק או בשיטת רבעים — לעולם לא \"חופן מלמעלה\"",
               "מסמנים: מקור, תאריך, נקודה ושם הדוגם"], 7.6, 1.2, 4.95, gap=0.76, size=14)
rows = [["Dmax (מ\"מ)", "4", "8", "16", "32", "63"], ["מסה מינ' (ק\"ג)", "0.2", "0.6", "2.6", "10", "40"]]
text(s, 0.9, 4.75, 6.4, 0.35, "מסה מינימלית מקובלת לניפוי", size=14, bold=True, color=NAVY)
table(s, 0.9, 5.15, 6.4, rows, [1.6, 0.96, 0.96, 0.96, 0.96, 0.96], size=13, row_h=0.42, header_h=0.42,
      align=["r", "c", "c", "c", "c", "c"], header_size=13)
message(s, "תוצאה מדויקת מדגימה לא מייצגת — חסרת ערך", x=7.6, w=4.95, y=5.1, h=0.7, ic="FaExclamationTriangle")

# 3 — grading & fineness modulus
s = content_slide("גרדציה ומודול דקות")
cd = CategoryChartData()
cd.categories = ["0.063", "0.125", "0.25", "0.5", "1", "2", "4", "8", "16", "22.4"]
cd.add_series(rl("גבול תחתון"), [0, 2, 6, 12, 20, 28, 36, 52, 78, 95])
cd.add_series(rl("גבול עליון"), [3, 8, 16, 26, 36, 46, 58, 76, 95, 100])
cd.add_series(rl("דגימה"), [1.5, 4, 10, 17, 25, 33, 44, 61, 86, 99])
gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, I(0.9), I(1.15), I(5.9), I(4.3), cd)
ch = gf.chart
style_chart(ch)
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.BOTTOM
ch.legend.include_in_layout = False
ser = ch.plots[0].series
series_line(ser[0], GRAY, w=2, marker=False, dash=True)
series_line(ser[1], GRAY, w=2, marker=False, dash=True)
series_line(ser[2], BLUE, w=3)
axis_fmt(ch.value_axis, rl("% עובר"))
ch.value_axis.minimum_scale, ch.value_axis.maximum_scale, ch.value_axis.major_unit = 0, 100, 20
axis_fmt(ch.category_axis, rl("פתח נפה (מ\"מ)"), gridlines=False)
caption(s, 0.9, 5.45, 5.9, "עקומת דירוג בתוך מעטפת הדרישה — דוגמה להמחשה", align="c")
text(s, 7.1, 1.15, 5.45, 0.4, "אופן הביצוע (עקום דירוג לפי ת\"י 3)", size=17, bold=True, color=NAVY)
steps_list(s, ["ייבוש ב-110±5°C ושקילה (M1)", "שטיפה על נפה 0.063 מ\"מ, ייבוש ושקילה (M2)",
               "ניפוי בסדרת נפות ושקילת השארית בכל נפה", "חישוב % עובר מצטבר והשוואה למעטפת"],
           7.1, 1.6, 5.45, gap=0.55, size=14, d=0.4)
box(s, 7.1, 3.9, 2.65, 1.75, fill=TINT)
text(s, 7.2, 3.95, 2.45, 1.65, [{"t": "מודול דקות FM", "bold": True, "color": NAVY, "size": 15},
                                {"t": "סכום ה-% הנשאר המצטבר על הנפות 0.15-4.75, חלקי 100", "size": 12},
                                {"t": "מקובל לחול בטון: 2.3-3.1", "bold": True, "color": BLUE, "size": 14}], align="c",
     anchor="m", space_after=2)
box(s, 9.9, 3.9, 2.65, 1.75, fill=TINT)
text(s, 10.0, 3.95, 2.45, 1.65, [{"t": "דקים מתחת ל-0.063", "bold": True, "color": NAVY, "size": 15},
                                 {"t": "ת\"י 3: תכולת אבק מרבית לפי סוג האגרגט (א' / ב')", "size": 12},
                                 {"t": "בחול גרוס, במהדורת 1998: 3% (א') / 5% (ב')", "size": 12}], align="c",
     anchor="m", space_after=2)
caption(s, 7.1, 5.75, 5.45, REQ_NOTE)

# 4 — fines quality: sand equivalent & methylene blue
s = content_slide("איכות הדקים: שווה חול וכחול מתילן")
pic_fit(s, "sand_eq", 0.9, 1.2, 4.3, 3.6)
text(s, 3.05, 4.85, 2.15, 0.6, [{"t": "חול חרסיתי", "bold": True, "color": RED}, {"t": "SE נמוך", "size": 12}],
     size=14, align="c", space_after=0)
text(s, 0.9, 4.85, 2.15, 0.6, [{"t": "חול נקי", "bold": True, "color": GREEN}, {"t": "SE גבוה", "size": 12}],
     size=14, align="c", space_after=0)
for k, (t1, std_, how, form, req) in enumerate((
        ("שווה חול (SE)", "שיטה מקובלת · ת\"י 3 / מפרט",
         ["חול 0/2 בתמיסה מפרידה בגליל מדורג", "ניעור ממוכן ושקיעה של 20 דקות", "מודדים את גובה החרסית (h1) ואת גובה החול (h2)"],
         "SE = h2 / h1 × 100", "טיפוסי במפרטים: SE ≥ 60-70"),
        ("כחול מתילן (MB)", "שיטה מקובלת · ת\"י 3 / מפרט",
         ["חלקיק 0/2 בתרחיף מים", "טיטרציה במנות של תמיסת כחול מתילן", "בדיקת \"הילה\" על נייר סינון עד רוויה"],
         "MB = גרם צבע לק\"ג חומר", "טיפוסי במפרטים: MB ≤ 1.0-2.0"))):
    x = 8.95 - k * 3.75
    box(s, x, 1.2, 3.6, 4.4, fill=TINT)
    text(s, x + 0.15, 1.3, 3.3, 0.45, t1, size=19, bold=True, color=NAVY)
    text(s, x + 0.15, 1.75, 3.3, 0.3, std_, size=12, color=MUTED)
    text(s, x + 0.15, 2.15, 3.3, 1.6, [{"t": it, "bullet": True} for it in how], size=13, space_after=3)
    r = box(s, x + 0.2, 3.85, 3.2, 0.6, fill=NAVY)
    shape_text(r, [[(form, {"size": 15})]])
    text(s, x + 0.15, 4.6, 3.3, 0.85, req, size=14, bold=True, color=BLUE, align="c", anchor="m")
message(s, "SE מסנן מהר; MB מבדיל בין אבק אבן תקין לחרסית מזיקה", y=5.8, h=0.62)
verify(s, 0.9, 0.55)

# 5 — density, absorption & moisture correction
s = content_slide("צפיפות, ספיגות ותיקון מים")
text(s, 7.1, 1.15, 5.45, 0.4, "אופן הביצוע (צפיפות וספיגות, ת\"י 3)", size=16, bold=True, color=NAVY)
steps_list(s, ["השריה במים 24 שעות", "ייבוש פני השטח עד מצב רווי-יבש-פנים (SSD)",
               "שקילה במצב SSD ושקילה בתוך מים (סל / פיקנומטר)", "ייבוש בתנור ושקילה יבשה"],
           7.1, 1.6, 5.45, gap=0.55, size=14, d=0.4)
box(s, 7.1, 3.9, 5.45, 0.75, fill=NAVY)
text(s, 7.2, 3.92, 5.25, 0.7, [{"t": "WA24 = (M_SSD − M_dry) / M_dry × 100", "rtl": False, "size": 17, "bold": True,
                                "color": WHITE}], align="c", anchor="m")
text(s, 7.1, 4.8, 5.45, 0.85, [{"t": "ת\"י 3: ספיגות מרבית 5% (לפי גיליון תיקון 2021)", "bullet": True, "bold": True, "color": BLUE},
                               {"t": "גיר/דולומיט איכותי: לרוב 0.5-2%; ערך גבוה — לבירור", "bullet": True}], size=14)
# worked example: moisture correction
box(s, 0.9, 1.2, 5.9, 4.45, fill=TINT2)
text(s, 1.1, 1.3, 5.5, 0.45, "דוגמה: תיקון לחות יומי", size=18, bold=True, color=NAVY)
rows = [["חומר", "כמות (ק\"ג)", "לחות", "ספיגות", "מים חופשיים"],
        ["חול", "800", "5.0%", "1.5%", "+28 ל'"],
        ["חצץ", "1,000", "0.5%", "1.2%", "−7 ל'"],
        ["סה\"כ", "", "", "", "+21 ל'"]]
table(s, 1.1, 1.85, 5.5, rows, [1.0, 1.2, 1.0, 1.0, 1.3], size=14, row_h=0.48, header_h=0.48,
      align=["r", "c", "c", "c", "c"], cell_fill=lambda i, j: LIME_T if i == 3 else None)
text(s, 1.1, 3.95, 5.5, 1.6, [{"t": "האגרגט מכניס 21 ליטר מים \"בשקט\" — מפחיתים אותם ממי התערובת", "bullet": True},
                              {"t": "ב-350 ק\"ג צמנט, בלי תיקון: W/C עולה בכ-0.06", "bullet": True,
                               "bold": True, "color": RED}], size=14)
caption(s, 0.9, 5.75, 11.65, "מים חופשיים = (לחות − ספיגות) × משקל יבש. הלחות נמדדת יומית (ייבוש דוגמה) או בחיישן במגורה")

# 6 — LA abrasion & particle shape
s = content_slide("עמידות לשחיקה וצורת גרגר")
pic_fit(s, "la_drum", 0.9, 1.2, 4.2, 3.2)
caption(s, 0.9, 4.45, 4.2, "תוף לוס אנג'לס: אגרגט + כדורי פלדה", align="c")
box(s, 5.4, 1.2, 7.15, 2.6, fill=TINT)
text(s, 5.6, 1.3, 6.75, 0.45, [[("שחיקה בתוף לוס אנג'לס (LA)", {"bold": True, "color": NAVY, "size": 19}),
                                ("   ת\"י 3", {"color": MUTED, "size": 12})]])
text(s, 5.6, 1.8, 4.1, 1.95, [{"t": "5 ק\"ג אגרגט 10-14 מ\"מ + 11 כדורי פלדה", "bullet": True},
                              {"t": "500 סיבובים ב-31-33 סל\"ד", "bullet": True},
                              {"t": "ניפוי על 1.6 מ\"מ; LA = % החומר שנשחק", "bullet": True}], size=13)
r = box(s, 9.85, 1.85, 2.5, 1.8, fill=WHITE, line=BLUE, lw=1.5)
shape_text(r, [[("דרישה", {"size": 14, "color": NAVY})], [("≤ 40", {"size": 26, "color": BLUE})],
               [("בטון מבני (טיפוסי)", {"size": 11, "color": TEXT, "bold": False})],
               [("לפי ת\"י 3 / מפרט", {"size": 11, "color": MUTED, "bold": False})]])
box(s, 5.4, 4.0, 7.15, 1.65, fill=TINT)
text(s, 5.6, 4.08, 6.75, 0.45, [[("גרגרים פחוסים ומוארכים", {"bold": True, "color": NAVY, "size": 19}),
                                 ("   ת\"י 3", {"color": MUTED, "size": 12})]])
text(s, 5.6, 4.55, 4.1, 1.05, [{"t": "מדידה במד צורה; % הגרגרים הפחוסים והמוארכים", "bullet": True},
                               {"t": "גרגר שטוח = יותר מים, פחות עבידות", "bullet": True}], size=13)
r = box(s, 9.85, 4.55, 2.5, 1.0, fill=WHITE, line=BLUE, lw=1.5)
shape_text(r, [[("תכולה מרבית", {"size": 16, "color": BLUE})],
               [("לפי ת\"י 3", {"size": 11, "color": TEXT, "bold": False})]])
caption(s, 0.9, 5.8, 11.65, REQ_NOTE)

# 7 — harmful substances
s = content_slide("חומרים מזיקים באגרגט")
harm = [("FaLeaf", "חומר אורגני", "ת\"י 3", "תמיסת NaOH והשוואת צבע לסטנדרט",
         "כהה מהסטנדרט ← בדיקת טיט השוואתית", "מעכב התקשרות"),
        ("FaCubes", "חרסית", "ת\"י 3", "תכולת חרסית וגושים פריכים",
         "תכולה מרבית לפי ת\"י 3", "חללים ונקודות חולשה"),
        ("FaFlask", "כלורידים", "ת\"י 118", "כלורידים מכל מרכיבי התערובת",
         "תקרה בבטון כולו, באחוז ממשקל הצמנט — לפי דרגת החשיפה", "קורוזיה של הזיון"),
        ("FaVials", "סולפטים וגופרית", "ת\"י 3", "סולפטים מסיסים וגופרית כללית",
         "תכולה מרבית לפי ת\"י 3", "התפחות והתפוררות"),
        ("FaBolt", "תגובה אלקלית (ASR)", "בדיקת מעבדה", "התפשחות מוטות טיט בתמיסה בסיסית",
         "התפשחות גבוהה = אגרגט תגובתי, לבירור", "סדיקה מתמשכת"),
        ("FaSearch", "חלקיקים קלים", "ת\"י 3", "הפרדה בנוזל כבד",
         "תכולה מרבית לפי ת\"י 3 / מפרט", "כתמים ובליטות בפני השטח")]
for k, (ic, t1, st, how, req, eff) in enumerate(harm):
    col_, row_ = k % 3, k // 3
    x = 12.55 - (col_ + 1) * 3.88 + 0.08
    y = 1.2 + row_ * 2.3
    box(s, x, y, 3.8, 2.15, fill=TINT)
    icon(s, ic, x + 3.4, y + 0.38, 0.5, fill=NAVY)
    text(s, x + 0.15, y + 0.12, 2.9, 0.4, t1, size=16, bold=True, color=NAVY)
    text(s, x + 0.15, y + 0.5, 2.9, 0.3, st, size=11, color=MUTED)
    text(s, x + 0.15, y + 0.85, 3.5, 0.4, how, size=12)
    text(s, x + 0.15, y + 1.22, 3.5, 0.55, req, size=12, bold=True, color=BLUE)
    text(s, x + 0.15, y + 1.78, 3.5, 0.3, "השפעה: " + eff, size=11, color=RED)
message(s, "מקור אגרגט חדש עובר את כל הבדיקות האלה לפני השימוש הראשון", y=5.85, h=0.6)

# 8 — lab report vs requirements
s = content_slide("דוח בדיקה מול דרישות: איך מחליטים")
OK_, NO_ = "עומד", "חורג"
rows = [["חומר", "בדיקה", "לפי", "תוצאה", "דרישה", "סטטוס"],
        ["חול גרוס", "עקום דירוג", "ת\"י 3", "בתוך התחום", "תחום התקן", OK_],
        ["חול גרוס", "תכולת אבק", "ת\"י 3", "7.8%", "≤ 5% (סוג ב')", NO_],
        ["חול גרוס", "שווה חול", "מפרט", "58", "≥ 65", NO_],
        ["חול גרוס", "כחול מתילן", "מפרט", "2.6 g/kg", "≤ 1.5 g/kg", NO_],
        ["חול גרוס", "ספיגות", "ת\"י 3", "1.4%", "≤ 5%", OK_],
        ["חצץ", "שחיקה LA", "ת\"י 3", "27", "לפי ת\"י 3", OK_],
        ["חצץ", "פחוסים ומוארכים", "ת\"י 3", "14%", "לפי ת\"י 3", OK_]]
table(s, 0.9, 1.2, 7.9, rows, [1.5, 1.45, 1.15, 1.35, 1.45, 1.0], size=13, row_h=0.47, header_h=0.47,
      align=["r", "r", "c", "c", "c", "c"],
      cell_fill=lambda i, j: (GREEN_T if rows[i][5] == OK_ else RED_T) if j == 5 else None)
caption(s, 0.9, 5.05, 7.9, "דוגמה להמחשה. " + REQ_NOTE)
verify(s, 0.9, 0.55)
box(s, 9.1, 1.2, 3.45, 4.3, fill=NAVY)
text(s, 9.3, 1.3, 3.05, 0.45, "המסקנה", size=19, bold=True, color=LIME)
text(s, 9.3, 1.8, 3.05, 1.2, "שלוש חריגות שמצביעות על אותה בעיה: הדקים בחול הם חרסית, לא אבק אבן", size=14,
     color=WHITE)
text(s, 9.3, 3.05, 3.05, 2.4, [{"t": "הסגר המשלוח וסימון הערמה", "bullet": True},
                               {"t": "דוח אי-התאמה והודעה לספק", "bullet": True},
                               {"t": "דרישה לשטיפה או להחלפת מקור", "bullet": True},
                               {"t": "בדיקה חוזרת לפני שחרור", "bullet": True}], size=13, color=WHITE)
message(s, "לא מחליטים לפי בדיקה אחת: מחפשים את הסיפור שהתוצאות מספרות יחד", y=5.7, h=0.7, ic="FaSearch")

# ---------------------------------------------------------------- 40. Cement tests
s = content_slide("בדיקות איכות צמנט")
cols_ = [("FaCertificate", "תעודות יצרן", ["תו תקן לפי ת\"י 1", "תעודת משלוח: סוג, דרגה ומשקל",
                                             "דוח איכות תקופתי מהיצרן"], "כל משלוח"),
         ("FaHourglassHalf", "זמן התקשרות", ["מחט ויקה — לפי ת\"י 1", "תחילת התקשרות: לפחות 45-75 דקות, לפי הדרגה",
                                              "חריגה מעידה על בעיית ייצור או אחסון"], "תקופתית / בחשד"),
         ("FaCompressArrowsAlt", "חוזק", ["מנסרות טיט 40×40×160 — לפי ת\"י 1", "בגיל 2 / 7 ו-28 יום",
                                          "השוואה לדרגה: 42.5 / 52.5 MPa"], "לפי תכנית")]
for i, (ic, t1, items, fr) in enumerate(cols_):
    x = 12.55 - (i + 1) * 3.88 + 0.08
    box(s, x, 1.2, 3.8, 3.9, fill=TINT)
    icon(s, ic, x + 1.9, 1.95, 1.0, fill=BLUE if i != 1 else NAVY)
    text(s, x + 0.15, 2.6, 3.5, 0.45, t1, size=20, bold=True, color=NAVY, align="c")
    text(s, x + 0.25, 3.15, 3.3, 1.5, [{"t": it, "bullet": True} for it in items], size=13)
    r = box(s, x + 0.8, 4.65, 2.2, 0.35, fill=LIME, radius=0.5)
    shape_text(r, fr, size=12, color=NAVY)
message(s, "בודקים גם את האחסון: צמנט חם או לח משנה התקשרות ועבידות", y=5.4, h=0.75)

# ---------------------------------------------------------------- 41. Admixture tests
s = content_slide("בדיקות איכות מוספים")
text(s, 0.9, 1.1, 11.65, 0.45, "בדיקות זהות מהירות בקבלה — השוואה לערך שהצהיר היצרן, לפי דרישות ת\"י 118 למרכיבי הבטון",
     size=15, color=MUTED)
adm = [("FaWeightHanging", "צפיפות", "±0.02-0.03", "גרם/סמ\"ק מהערך המוצהר", "הידרומטר / פיקנומטר"),
       ("FaFlask", "pH", "±1", "יחידת pH מהמוצהר", "מד pH מכויל"),
       ("FaFilter", "מוצקים", "±5%", "יחסי מהערך המוצהר", "ייבוש ב-105°C"),
       ("FaTint", "צמיגות", "השוואה", "לדגימת ייחוס מאושרת", "צמיגומטר / משפך")]
for i, (ic, t1, v, unit, how) in enumerate(adm):
    x = 12.55 - (i + 1) * 2.92 + 0.07
    box(s, x, 1.75, 2.78, 3.6, fill=TINT)
    icon(s, ic, x + 1.39, 2.45, 0.9, fill=BLUE if i % 2 == 0 else NAVY)
    text(s, x + 0.1, 3.05, 2.58, 0.45, t1, size=20, bold=True, color=NAVY, align="c")
    text(s, x + 0.1, 3.55, 2.58, 0.55, v, size=26, bold=True, color=BLUE, align="c", rtl=v == "השוואה")
    text(s, x + 0.1, 4.12, 2.58, 0.4, unit, size=12, align="c")
    text(s, x + 0.1, 4.7, 2.58, 0.4, how, size=12, color=MUTED, align="c")
message(s, "לבדוק גם: תאריך תפוגה, הפרדת פאזות במיכל וזיהוי נכון של המוסף בצנרת המינון", y=5.6, h=0.75)

# ---------------------------------------------------------------- 42. Rejected materials gallery
s = content_slide("זיהוי חומרים פסולים")
tiles = [("tile_ok", True, "תקין", "אגרגט נקי, צבע אחיד, ללא ציפוי"),
         ("tile_clay", False, "חרסית", "ציפוי וגושים חומים; נמרח ביד רטובה"),
         ("tile_lumps", False, "גושי צמנט", "גושים קשים שאינם מתפוררים באצבע"),
         ("tile_organic", False, "חומר אורגני", "שורשים, עלים, צבע כהה וריח"),
         ("tile_phase", False, "הפרדת פאזות", "שכבות נפרדות במיכל המוסף")]
tw_, gap_ = 2.2, 0.16
for i, (img, ok, t1, t2) in enumerate(tiles):
    x = 12.55 - (i + 1) * tw_ - i * gap_
    col = GREEN if ok else RED
    box(s, x, 1.25, tw_, 4.4, fill=GREEN_T if ok else RED_T)
    pic_fit(s, img, x + 0.1, 1.35, tw_ - 0.2, 1.6)
    icon(s, "FaCheckCircle" if ok else "FaTimesCircle", x + tw_ / 2, 3.35, 0.65, fill=col)
    text(s, x + 0.1, 3.8, tw_ - 0.2, 0.45, t1, size=19, bold=True, color=col, align="c")
    text(s, x + 0.15, 4.3, tw_ - 0.3, 1.2, t2, size=13, align="c")
caption(s, 0.9, 5.75, 11.65, "איורים להמחשה. מומלץ להחליף בתמונות אמיתיות מהמפעלים, לצורך זיהוי מדויק בשטח", align="c")

# ---------------------------------------------------------------- 43. Risk matrix
s = content_slide("מטריצת סיכוני איכות")
risks = [("אגרגט דק", "חרסית / עודף דקים", 3, 4), ("אגרגט גס", "גרדציה חורגת", 3, 3),
         ("צמנט", "סוג שגוי בממגורה", 2, 5), ("צמנט", "לחות וגושים", 2, 3),
         ("מוסף", "מינון שגוי", 2, 5), ("מוסף", "הפרדת פאזות", 2, 4), ("מים", "מים ממוחזרים לא מבוקרים", 2, 3)]


def lvl(v):
    return ("גבוה", RED_T) if v >= 10 else ("בינוני", AMBER_T) if v >= 6 else ("נמוך", GREEN_T)


rows = [["#", "חומר גלם", "סיכון", "הסתברות", "השפעה", "ציון", "רמה"]]
for k, (m_, r_, p_, i_) in enumerate(risks):
    rows.append([str(k + 1), m_, r_, str(p_), str(i_), str(p_ * i_), lvl(p_ * i_)[0]])
table(s, 0.9, 1.2, 8.4, rows, [0.45, 1.35, 2.75, 1.05, 0.95, 0.75, 1.0], size=13, row_h=0.5,
      align=["c", "r", "r", "c", "c", "c", "c"], bold_first=True,
      cell_fill=lambda i, j: lvl(risks[i - 1][2] * risks[i - 1][3])[1] if j in (5, 6) else None)
caption(s, 0.9, 5.3, 8.4, "סולם 1-5. דוגמה — יש לכייל את הציונים לפי נתוני המפעל ולעדכן אחת לרבעון")
# heat-map: impact grows to the left (RTL), probability grows upwards
hx, hy, cs = 9.75, 1.65, 0.52
text(s, hx, 1.2, 5 * cs, 0.35, "הסתברות ↑   השפעה ←", size=12, color=MUTED, align="c")
for pi in range(5):
    for ii in range(5):
        v = (5 - pi) * (ii + 1)
        col = "F4B7B7" if v >= 10 else "FFD99A" if v >= 6 else "BFE3C2"
        box(s, hx + (4 - ii) * cs, hy + pi * cs, cs - 0.03, cs - 0.03, fill=col, shape=MSO_SHAPE.RECTANGLE)
cells = {}
for k, (m_, r_, p_, i_) in enumerate(risks):
    cells.setdefault((p_, i_), []).append(k + 1)
for (p_, i_), ids in cells.items():
    x0_ = hx + (5 - i_) * cs
    y0_ = hy + (5 - p_) * cs
    for j, n in enumerate(ids):
        off = (j - (len(ids) - 1) / 2) * 0.2
        badge(s, n, x0_ + cs / 2 - 0.015 + off, y0_ + cs / 2 - 0.015, 0.24, fill=NAVY, size=9)
text(s, 9.55, 4.4, 3.0, 0.7, "מתחילים מהתאים האדומים: שם נמצא רוב הסיכון", size=13, color=NAVY, bold=True,
     align="c")

# ================================================================ Section 9 — recycled materials
divider(9, "חומרים ממוחזרים", ["מים ממוחזרים", "בטון חוזר", "אגרגט ממוחזר", "מתי מותר"])

# R1 — overview
s = content_slide("חומרים ממוחזרים במפעל הבטון")
streams = [("FaWater", "מים ממוחזרים", "מי שטיפת משאיות ומפעל אחרי שיקוע", "כמעט בכל מפעל", BLUE),
           ("FaSyncAlt", "בטון חוזר", "עודפי בטון טרי שחוזרים מהאתר", "אחוזים בודדים מהייצור", NAVY),
           ("FaRecycle", "אגרגט ממוחזר", "גריסת בטון קשוח ופסולת בניין", "בהיקף גדל", LIME_D)]
for i, (ic, t1, t2, t3, col) in enumerate(streams):
    x = 12.55 - (i + 1) * 3.88 + 0.08
    box(s, x, 1.2, 3.8, 2.4, fill=TINT)
    icon(s, ic, x + 1.9, 1.85, 0.9, fill=col)
    text(s, x + 0.15, 2.4, 3.5, 0.45, t1, size=20, bold=True, color=NAVY, align="c")
    text(s, x + 0.2, 2.85, 3.4, 0.4, t2, size=13, align="c")
    text(s, x + 0.2, 3.2, 3.4, 0.35, t3, size=12, color=MUTED, align="c")
box(s, 6.4, 3.8, 6.15, 1.75, fill=GREEN_T)
icon(s, "FaLeaf", 12.1, 4.25, 0.5, fill=GREEN)
text(s, 6.6, 3.95, 5.0, 0.4, "למה ממחזרים", size=17, bold=True, color=GREEN)
text(s, 6.6, 4.4, 5.7, 1.1, [{"t": "חיסכון במים ובחומרי גלם; פחות פסולת להטמנה", "bullet": True},
                             {"t": "רגולציה סביבתית ודרישות בנייה ירוקה", "bullet": True}], size=14)
box(s, 0.9, 3.8, 5.3, 1.75, fill=AMBER_T)
icon(s, "FaExclamationTriangle", 5.75, 4.25, 0.5, fill=AMBER)
text(s, 1.1, 3.95, 4.2, 0.4, "מה הסיכון", size=17, bold=True, color=AMBER)
text(s, 1.1, 4.4, 4.9, 1.1, [{"t": "שונות גבוהה בין יום ליום", "bullet": True},
                             {"t": "מוצקים, מעכבים, ספיגות וזיהומים", "bullet": True}], size=14)
message(s, "ממוחזר זה לא פסול: זה חומר גלם שדורש בקרה צמודה יותר", y=5.75, h=0.65)

# R2 — recycled water system
s = content_slide("מים ממוחזרים: מקורות ומערכת")
wa = json.load(open(f"{ASSETS}/water_anchors.json"))
wx, wy, ww = 0.9, 1.2, 11.65
pic(s, "water_plant", wx, wy, w=ww)
wsc = ww / 1400
labels = ["שטיפת המשאית", "מפריד (Reclaimer): חצץ וחול חוזרים", "בריכות שיקוע", "מיכל ערבול + מד צפיפות",
          "מינון במפעל"]
for k, (ax, lb) in enumerate(zip(wa, labels)):
    cx_ = wx + ax * wsc
    badge(s, k + 1, cx_, 4.55, 0.42, fill=NAVY, size=13)
    text(s, cx_ - 1.15, 4.8, 2.3, 0.6, lb, size=13, align="c", bold=True, color=NAVY)
cards = [("FaCheckCircle", GREEN, GREEN_T, "מקורות מקובלים", "מי שטיפה, ניקוז המפעל, מי גשם מהחצר"),
         ("FaBan", RED, RED_T, "לא נכנס למערכת", "שמנים, דלק, דטרגנטים, שפכים סניטריים"),
         ("FaSyncAlt", BLUE, TINT, "ערבול רציף", "בלי ערבול המוצקים שוקעים והצפיפות \"קופצת\"")]
# three compact cards in one row, kept left of the logo zone
for i, (ic, col, bg, t1, t2) in enumerate(cards):
    x = 6.9 - i * 3.0
    box(s, x, 5.45, 2.9, 1.1, fill=bg)
    icon(s, ic, x + 2.58, 5.75, 0.42, fill=col)
    text(s, x + 0.12, 5.53, 2.2, 0.35, t1, size=14, bold=True, color=col)
    text(s, x + 0.12, 5.9, 2.25, 0.62, t2, size=12)

# R3 — mixing-water requirements
s = content_slide("מים ממוחזרים: בדיקות ודרישות לפי ת\"י 118")
rows = [["שלב", "פרמטר", "דרישה", "איך ומתי"],
        ["בדיקה מקדימה", "שמנים ושומנים", "לא יותר מעקבות נראות", "חזותית, יומי"],
        ["בדיקה מקדימה", "דטרגנטים", "קצף נעלם תוך 2 דקות", "ניעור בבקבוק, יומי"],
        ["בדיקה מקדימה", "ריח ו-pH", "ללא ריח חריג; pH ≥ 4", "יומי / שבועי"],
        ["כימיה", "כלורידים (Cl)", "≤ 500 דרוך · ≤ 1,000 מזוין · ≤ 4,500 לא מזוין", "מעבדה, לפי תכנית"],
        ["כימיה", "סולפטים (SO₄)", "≤ 2,000", "מעבדה, לפי תכנית"],
        ["כימיה", "אלקליים (Na₂O שקול)", "≤ 1,500", "מעבדה, לפי תכנית"],
        ["כימיה", "סוכרים · פוספטים · חנקות", "≤ 100 · ≤ 100 · ≤ 500", "מעבדה, בחשד"],
        ["כימיה", "עופרת · אבץ", "≤ 100 · ≤ 100", "מעבדה, בחשד"]]
table(s, 0.9, 1.2, 11.65, rows, [1.8, 2.75, 4.4, 2.7], size=14, row_h=0.49, header_h=0.48,
      align=["r", "r", "c", "r"], cell_fill=lambda i, j: (TINT if i <= 3 else LIME_T) if j == 0 else None)
caption(s, 0.9, 5.65, 11.65, "ריכוזים ב-mg/l. ת\"י 118 קובע את הדרישות למי ערבוב, כולל מים ממוחזרים מתעשיית הבטון")
verify(s, 0.9, 0.55)
caption(s, 0.9, 5.95, 9.0, "ת\"י 118 מפנה לעניין זה לתקן המים האירופי; יש לאמת את הערכים מול המהדורה המחייבת")

# R4 — performance tests & solids via density
s = content_slide("מים ממוחזרים: התקשרות, חוזק ומוצקים")
box(s, 7.1, 1.2, 5.45, 2.4, fill=TINT)
text(s, 7.3, 1.3, 5.05, 0.4, "בדיקת ביצועים מול מים מזוקקים", size=17, bold=True, color=NAVY)
text(s, 7.3, 1.8, 5.05, 1.75, [{"t": "התחלת התקשרות ≥ 1 שעה; סוף ≤ 12 שעות", "bullet": True},
                               {"t": "סטייה מזמני הביקורת: עד 25%", "bullet": True},
                               {"t": "חוזק 7 ימים ≥ 90% מהביקורת", "bullet": True, "bold": True, "color": BLUE}], size=14)
verify(s, 10.6, 3.2)
box(s, 7.1, 3.8, 5.45, 1.75, fill=NAVY)
text(s, 7.3, 3.85, 5.05, 1.65, [{"t": "כלל המוצקים", "size": 15, "bold": True, "color": LIME},
                                {"t": "מסת המוצקים מהמים הממוחזרים ≤ 1% ממסת האגרגט הכוללת", "size": 15,
                                 "color": WHITE},
                                {"t": "דוגמה: 170 ל' מים ו-1,900 ק\"ג אגרגט — עד 19 ק\"ג מוצקים, כלומר צפיפות מים של 1.06 לכל היותר", "size": 12,
                                 "color": WHITE}], anchor="m")
dens = [1.01, 1.02, 1.03, 1.04, 1.05, 1.06, 1.08, 1.10]
cd = CategoryChartData()
cd.categories = [f"{d_:.2f}" for d_ in dens]
cd.add_series(rl("מוצקים (גרם לליטר)"), [round((d_ - 1) / (1 - 1 / 2.1) * 1000) for d_ in dens])
cd.add_series(rl("גבול לתערובת לדוגמה"), [112] * len(dens))
gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, I(0.9), I(1.15), I(5.95), I(4.35), cd)
ch = gf.chart
style_chart(ch)
ch.has_legend = True
ch.legend.position = XL_LEGEND_POSITION.BOTTOM
ch.legend.include_in_layout = False
series_line(ch.plots[0].series[0], BLUE, w=3)
series_line(ch.plots[0].series[1], RED, w=2, marker=False, dash=True)
axis_fmt(ch.value_axis, rl("גרם מוצקים לליטר"))
ch.value_axis.minimum_scale, ch.value_axis.maximum_scale, ch.value_axis.major_unit = 0, 200, 50
axis_fmt(ch.category_axis, rl("צפיפות המים (ק\"ג/ל')"), gridlines=False, fmt="0.00")
caption(s, 0.9, 5.55, 5.95, "חישוב בהנחת צפיפות מוצקים 2.1. המדידה היומית של הצפיפות היא הבקרה החשובה ביותר", align="c")

# R5 — usage rules
s = content_slide("כללי שימוש במים ממוחזרים")
dos = ["ערבול רציף במיכל לפני המינון", "מדידת צפיפות לפחות יומית (עדיף חיישן רציף)",
       "חישוב המוצקים והכללתם בתכן כחומר דק", "המים הממוחזרים נספרים במלואם ב-W/C",
       "כימיה, התקשרות וחוזק לפי תכנית הבקרה", "תיעוד שיעור המים הממוחזרים בכל מנה"]
cau = ["בטון אדריכלי חשוף — שינויי גוון", "בטון דרוך — מגבלת כלורידים מחמירה",
       "בטון עם אוויר מוחדר — השפעה על כמות האוויר", "חוזק גבוה ו-SCC — רגישות לדקים",
       "שאריות מייצב הידרציה — עיכוב התקשרות", "צפיפות מעל הגבול — מדללים במים נקיים"]
for k, (hdr, items, col, bg, ic) in enumerate((("עושים תמיד", dos, GREEN, GREEN_T, "FaCheckCircle"),
                                                ("זהירות מיוחדת", cau, AMBER, AMBER_T, "FaExclamationTriangle"))):
    x = 6.85 - k * 5.95
    box(s, x, 1.2, 5.7, 4.35, fill=bg)
    icon(s, ic, x + 5.2, 1.7, 0.6, fill=col)
    text(s, x + 0.25, 1.45, 4.5, 0.45, hdr, size=20, bold=True, color=col)
    text(s, x + 0.25, 2.2, 5.2, 3.3, [{"t": it, "bullet": True} for it in items], size=15, space_after=5)
message(s, "המים הממוחזרים הם חומר גלם לכל דבר: נבדקים, מתועדים ונכנסים לתכן", y=5.75, h=0.65)

# R6 — returned concrete: three routes
s = content_slide("בטון חוזר: שלושה מסלולים")
routes = [("FaSyncAlt", "שטיפה במפריד", "בטון טרי → Reclaimer → חצץ וחול שטופים + מים ממוחזרים",
           "אגרגט כמעט כמו טבעי; בודקים דקים וגרדציה", GREEN_T, GREEN),
          ("FaHourglassHalf", "ייצוב ושימוש חוזר", "מוסף מייצב הידרציה עוצר את ההתקשרות לשעות",
           "רק בהנחיית טכנולוג; לרוב לאלמנטים לא קריטיים", AMBER_T, AMBER),
          ("FaCubes", "התקשות וגריסה", "בטון קשוח → גריסה וניפוי → אגרגט ממוחזר",
           "אגרגט ממוחזר — רק במפרט מיוחד ובאישור מתכנן", TINT, NAVY)]
for i, (ic, t1, t2, t3, bg, col) in enumerate(routes):
    x = 12.55 - (i + 1) * 3.88 + 0.08
    box(s, x, 1.2, 3.8, 3.75, fill=bg)
    icon(s, ic, x + 1.9, 1.85, 0.9, fill=col)
    text(s, x + 0.15, 2.45, 3.5, 0.45, t1, size=19, bold=True, color=col, align="c")
    text(s, x + 0.2, 2.95, 3.4, 0.95, t2, size=13, align="c")
    text(s, x + 0.2, 3.95, 3.4, 0.9, t3, size=13, bold=True, color=NAVY, align="c")
message(s, "בטון חוזר מתועד במקורו: מאיזו תערובת, כמה זמן, ואם קיבל מוסף", y=5.3, h=0.7, ic="FaClipboardCheck")

# R7 — recycled aggregate: properties & tests
s = content_slide("אגרגט ממוחזר: תכונות ובדיקות")
pic_fit(s, "rca", 0.9, 1.2, 5.5, 2.3)
text(s, 3.8, 3.5, 2.6, 0.6, [{"t": "אגרגט ממוחזר", "bold": True, "color": NAVY},
                             {"t": "טיט דבוק נקבובי סביב הסלע", "size": 12}], size=14, align="c", space_after=0)
text(s, 0.9, 3.5, 2.6, 0.6, [{"t": "אגרגט טבעי", "bold": True, "color": NAVY},
                             {"t": "סלע נקי וצפוף", "size": 12}], size=14, align="c", space_after=0)
rows = [["תכונה", "טבעי", "ממוחזר"], ["צפיפות (ק\"ג/ל')", "2.6-2.7", "2.1-2.5"], ["ספיגות", "0.5-2%", "3-10%"],
        ["LA", "20-30", "30-45"], ["שונות", "נמוכה", "גבוהה"]]
table(s, 0.9, 4.2, 5.5, rows, [2.1, 1.7, 1.7], size=13, row_h=0.36, header_h=0.38, align=["r", "c", "c"])
text(s, 6.8, 1.2, 5.75, 0.45, "בדיקות ייעודיות", size=19, bold=True, color=NAVY)
tests = [("FaListOl", "מיון מרכיבים", "בטון, אבן טבעית, לבנים, אספלט, זיהומים וחומר צף"),
         ("FaTint", "ספיגות", "בסיס לתיקון מים; מרטיבים מראש את הערמה"),
         ("FaVials", "סולפטים מסיסים", "גבס מטיח ולוחות — סכנת התפחות"),
         ("FaFlask", "כלורידים", "בטון ישן מסביבה ימית או מחניונים"),
         ("FaTrashAlt", "זיהומים", "עץ, פלסטיק, זכוכית, מתכת, אספלט")]
for i, (ic, t1, t2) in enumerate(tests):
    y = 1.75 + i * 0.8
    icon(s, ic, 12.25, y + 0.32, 0.5, fill=BLUE)
    text(s, 6.8, y, 5.1, 0.35, t1, size=15, bold=True, color=NAVY)
    text(s, 6.8, y + 0.33, 5.1, 0.45, t2, size=12)

# R8 — recycled aggregate in Israeli standards
s = content_slide("אגרגט ממוחזר בתקינה הישראלית")
stat = [("FaBook", "ת\"י 3", "חל על אגרגטים ממקור טבעי ומגריסה — לא על אגרגט ממוחזר", NAVY),
        ("FaBook", "ת\"י 118", "במהדורה המעודכנת: נספח למידע בלבד עם המלצות לשימוש באגרגט ממוחזר גס", BLUE),
        ("FaUserCheck", "המשמעות למפעל", "שימוש רק במפרט מיוחד, באישור המתכנן ובבקרה מוגברת", LIME_D)]
for i, (ic, t1, t2, col) in enumerate(stat):
    x = 12.55 - (i + 1) * 3.88 + 0.08
    box(s, x, 1.2, 3.8, 2.45, fill=TINT)
    icon(s, ic, x + 1.9, 1.8, 0.8, fill=col)
    text(s, x + 0.15, 2.3, 3.5, 0.45, t1, size=20, bold=True, color=NAVY, align="c")
    text(s, x + 0.2, 2.75, 3.4, 0.85, t2, size=13, align="c")
verify(s, 4.7, 3.75, w=4.0, t="מעמד הנספח ונוסחו — לאימות במהדורה המחייבת")
text(s, 0.9, 4.2, 11.65, 0.4, "הנחיות מקצועיות למפעל", size=17, bold=True, color=NAVY)
guide = ["רק אגרגט גס; אגרגט ממוחזר דק — לא מומלץ", "מקור ידוע: עדיף בטון חוזר של המפעל עצמו",
         "בדיקות כמו לאגרגט טבעי + ספיגות, סולפטים וכלורידים", "לא בדרגות החשיפה החמורות, לא בדרוך ולא באדריכלי"]
for k, g in enumerate(guide):
    col_, row_ = k % 2, k // 2
    x = 6.85 - col_ * 5.95
    y = 4.65 + row_ * 0.55
    text(s, x, y, 5.7, 0.5, [{"t": g, "bullet": True}], size=14)
message(s, "בספק — מתייחסים לאגרגט ממוחזר כאל חומר לא מאושר עד אישור כתוב", y=5.85, h=0.6, ic="FaExclamationTriangle")

# R9 — decision matrix
s = content_slide("מטריצת החלטה: מתי מותר להשתמש")
YES, COND, NO = GREEN_T, AMBER_T, RED_T
mx = [["חומר", "בטון רזה / לא מזוין", "מזוין, דרגות 1-3", "דרגות ימיות וכימיות", "דרוך / טרום דרוך", "אדריכלי חשוף",
       "חוזק גבוה / SCC"],
      ["מים ממוחזרים (מוצקים ≤ 1%)", "מותר", "מותר", "בתנאים: כלורידים", "בתנאים: כלורידים", "לא מומלץ", "בתנאים"],
      ["אגרגט שטוף ממפריד", "מותר", "מותר", "בתנאים", "בתנאים", "בתנאים", "בתנאים"],
      ["בטון חוזר מיוצב", "בתנאים", "בתנאים", "לא", "לא", "לא", "לא"],
      ["אגרגט ממוחזר מבטון של המפעל", "בתנאים", "בתנאים", "לא", "לא", "לא", "לא"],
      ["אגרגט ממוחזר מפסולת בניין", "בתנאים", "לא", "לא", "לא", "לא", "לא"]]


def mcol(i, j):
    if j == 0:
        return TINT
    v = mx[i][j]
    return NO if v.startswith("לא") else YES if v == "מותר" else COND


table(s, 0.9, 1.2, 11.65, mx, [2.65, 1.45, 1.55, 1.75, 1.55, 1.25, 1.45], size=12, row_h=0.55, header_h=0.62,
      align=["r"] + ["c"] * 6, cell_fill=mcol, header_size=12)
for k, (lb, col) in enumerate((("מותר בבקרה רגילה", YES), ("בתנאים / עם מגבלה", COND), ("לא מומלץ / אסור", NO))):
    x = 12.55 - (k + 1) * 2.6
    box(s, x + 2.15, 4.85, 0.35, 0.3, fill=col, line=GRAY, lw=0.75, shape=MSO_SHAPE.RECTANGLE)
    text(s, x, 4.83, 2.05, 0.35, lb, size=13)
caption(s, 0.9, 5.3, 11.65, "המלצה מקצועית כללית — ת\"י 118, ת\"י 3, המפרט והמתכנן גוברים תמיד. דרגות החשיפה לפי ת\"י 118")
message(s, "בספק — לא משתמשים. ממוחזר נכנס רק עם בדיקה, תיעוד ואישור טכנולוג", y=5.7, h=0.7, ic="FaUserCheck")

# ================================================================ Section 10
divider(10, "חידון, תקנים וסיכום", ["5 שאלות לבדיקת הבנה", "מפת התקנים", "מסרים מרכזיים"])

quiz = [("מהו החומר הקושר בבטון?", ["אגרגט", "צמנט", "מים", "מוסף"]),
        ("כמה אחוז מנפח הבטון הם אגרגטים?", ["10%-20%", "30%-40%", "50%-60%", "70%-80%"]),
        ("מהו C-S-H?", ["סוג של מוסף", "תוצר ההידרציה העיקרי, שאחראי לחוזק", "סוג של צמנט", "בדיקת חוזק"]),
        ("מה קורה כאשר יחס מים-צמנט עולה?", ["החוזק עולה", "החדירות יורדת", "החוזק יורד והחדירות עולה",
                                               "אין השפעה"]),
        ("איזו בדיקה מבוצעת בבטון טרי?", ["קרוטים", "שקיעה", "LA Abrasion", "לחיצת גלילים בגיל 28"])]
for qi, (q, opts) in enumerate(quiz):
    s = content_slide(f"חידון {qi + 1}")
    box(s, 0.9, 1.25, 11.65, 1.55, fill=NAVY)
    icon(s, "FaQuestion", 11.75, 2.02, 1.0, fill=LIME, color="n")
    text(s, 1.2, 1.3, 9.9, 1.45, q, size=30, bold=True, color=WHITE, anchor="m")
    for k, o in enumerate(opts):
        col_, row_ = k % 2, k // 2
        x = 6.85 - col_ * 5.95
        y = 3.1 + row_ * 1.35
        box(s, x, y, 5.7, 1.1, fill=TINT)
        badge(s, "אבגד"[k], x + 5.2, y + 0.55, 0.65, fill=BLUE, size=20)
        text(s, x + 0.25, y, 4.5, 1.1, o, size=20, color=NAVY, anchor="m")

# ---------------------------------------------------------------- Standards overview (Israeli only)
s = content_slide("התקנים הישראליים לבטון")
il = [("FaWarehouse", "ת\"י 1", "צמנט", "סוגים (CEM I-V), דרגות חוזק ובדיקות"),
      ("FaCubes", "ת\"י 3", "אגרגטים לבטון", "עקום דירוג, אבק, חרסית, ספיגות, צורה"),
      ("FaFlask", "ת\"י 26", "שיטות בדיקה לבטון", "דגימה, בטון טרי, חוזק, קרוטים, בדיקות לא הורסות"),
      ("FaBook", "ת\"י 118", "בטון: דרישות, תפקוד וייצור", "דרגות חוזק וחשיפה, מרכיבים, בקרת ייצור והתאמה"),
      ("FaBuilding", "ת\"י 466", "חוקת הבטון", "תכן מבני; כיסוי ועמידות"),
      ("FaHardHat", "ת\"י 467", "ביצוע עבודות בטון", "יציקה, הידוק, אשפרה ופיקוח")]
for k, (ic, num, t1, t2) in enumerate(il):
    col_, row_ = k % 3, k // 3
    x = 12.55 - (col_ + 1) * 3.88 + 0.08
    y = 1.2 + row_ * 2.2
    box(s, x, y, 3.8, 2.05, fill=NAVY if num == "ת\"י 118" else TINT)
    dark = num == "ת\"י 118"
    icon(s, ic, x + 3.3, y + 0.5, 0.65, fill=LIME if dark else BLUE, color="n" if dark else "w")
    text(s, x + 0.2, y + 0.2, 2.6, 0.5, num, size=24, bold=True, color=LIME if dark else BLUE)
    text(s, x + 0.2, y + 0.75, 3.4, 0.45, t1, size=16, bold=True, color=WHITE if dark else NAVY)
    text(s, x + 0.2, y + 1.22, 3.4, 0.75, t2, size=13, color=WHITE if dark else TEXT)
message(s, "ת\"י 118 הוא התקן המרכזי: מגדיר את הדרישות ומפנה לתקני המרכיבים והבדיקות",
        y=5.75, h=0.7, ic="FaBook")

# ---------------------------------------------------------------- Standards appendix
std_il = [["תקן", "נושא", "מטרה", "תחום", "רלוונטי ל-"],
          ["ת\"י 1", "צמנט", "הרכב, סוגים ודרגות חוזק", "ייצור ואספקת צמנט", "איכות, רכש, מעבדה"],
          ["ת\"י 3", "אגרגטים לבטון", "דרישות לאגרגט טבעי וגרוס", "מחצבות ומפעלי אגרגט", "מחצבות, איכות, מעבדה"],
          ["ת\"י 26 ח' 1", "דגימה והכנת דגימות", "נטילת מדגמים מבטון טרי", "אתר, מפעל ומעבדה", "טכנאים, מעבדה"],
          ["ת\"י 26 ח' 2", "בדיקות בטון טרי", "סומך, צפיפות, אוויר ועוד", "אתר ומפעל", "טכנאים, איכות"],
          ["ת\"י 26 ח' 4", "חוזק בטון קשוי", "חוזק לחיצה של קוביות", "מעבדות", "מעבדה, טכנולוגים"],
          ["ת\"י 26 ח' 5", "בדיקות מיוחדות", "ספיגות, חדירת מים ועוד", "מעבדות", "מעבדה, מהנדסים"],
          ["ת\"י 26 ח' 6, 7", "קרוטים ובדיקות לא הורסות", "הערכת הבטון במבנה", "מבנים קיימים וכשלים", "מהנדסים, איכות"],
          ["ת\"י 118", "בטון: דרישות, תפקוד וייצור", "דרגות חוזק וחשיפה, התאמה", "כל בטון: באתר, טרום, מובא", "טכנולוגים, מנהלי מפעלים"],
          ["ת\"י 466", "חוקת הבטון", "תכן מבני ועמידות", "מבני בטון", "מהנדסים, מנהלי אזורים"],
          ["ת\"י 467", "ביצוע עבודות בטון", "יציקה, הידוק ואשפרה", "אתרי בנייה", "מנהלי אזורים, פיקוח"]]
s = content_slide("נספח: התקנים הישראליים הרלוונטיים")
table(s, 0.9, 1.15, 11.65, std_il, [1.65, 2.5, 2.9, 2.3, 2.3], size=12, row_h=0.42, header_h=0.46,
      cell_fill=lambda i, j: TINT if j == 0 else None)
caption(s, 0.9, 6.0, 9.0, "יש לוודא תמיד את המהדורה וגיליונות התיקון העדכניים באתר מכון התקנים הישראלי")

# ---------------------------------------------------------------- Summary
s = content_slide("סיכום: 5 מסרים מרכזיים")
msgs = [("FaCubes", "איכות מתחילה בחומרי הגלם"), ("FaTint", "W/C הוא הפרמטר הקריטי ביותר"),
        ("FaClipboardCheck", "בקרת איכות מונעת כשלים"), ("FaDollarSign", "זיהוי מוקדם חוסך כסף"),
        ("FaUsers", "איכות היא אחריות של כולם")]
for i, (ic, t1) in enumerate(msgs):
    y = 1.2 + i * 0.95
    box(s, 2.5, y, 10.05, 0.8, fill=NAVY if i == 1 else TINT)
    badge(s, i + 1, 12.05, y + 0.4, 0.58, fill=LIME, color=NAVY, size=20)
    icon(s, ic, 3.05, y + 0.4, 0.58, fill=BLUE)
    text(s, 3.6, y, 7.9, 0.8, t1, size=22, bold=True, color=WHITE if i == 1 else NAVY, anchor="m")
icon(s, "FaGraduationCap", 1.45, 3.1, 1.1, fill=NAVY)

# ---------------------------------------------------------------- Thanks
s = clone(TPL[7])
t = find(s, "TextBox 4")
t.left, t.width = I(3.67), I(6.0)
retext(t, "תודה", size=32, align="ctr")
text(s, 3.67, 4.35, 6.0, 0.6, "שאלות?", size=26, color=NAVY, align="c")

# ---------------------------------------------------------------- drop the template's sample slides
sldIdLst = prs.slides._sldIdLst
for sid in list(sldIdLst)[:8]:
    prs.part.drop_rel(sid.get(qn("r:id")))
    sldIdLst.remove(sid)

prs.save(OUT)
print("slides:", len(prs.slides), "->", OUT)
